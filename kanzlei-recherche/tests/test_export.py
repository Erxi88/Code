from datetime import date

import pytest
from openpyxl import load_workbook

from recherche.export import (
    BLATT_DUBLETTEN,
    BLATT_KANZLEIEN,
    BLATT_NICHT_GEFUNDEN,
    BLATT_PROTOKOLL,
    DROPDOWN_MIN_ZEILEN,
    KANZLEI_SPALTEN,
    export_excel,
)  # fmt: skip
from recherche.storage import Storage


@pytest.fixture
def wb(befuellt, tmp_path):
    pfad = export_excel(befuellt, "SH", tmp_path / "out" / "kanzleien.xlsx")
    assert pfad.exists()
    return load_workbook(pfad)


def _kopf(ws):
    return [c.value for c in ws[1]]


def _zeilen(ws):
    kopf = _kopf(ws)
    return [dict(zip(kopf, (c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=2)]


def test_vier_blaetter(wb):
    assert wb.sheetnames == [BLATT_KANZLEIEN, BLATT_NICHT_GEFUNDEN, BLATT_DUBLETTEN, BLATT_PROTOKOLL]


def test_kanzleien_kopf_und_zeilen(wb):
    ws = wb[BLATT_KANZLEIEN]
    kopf = _kopf(ws)
    namen = [s[0] for s in KANZLEI_SPALTEN]
    assert kopf == [namen[0], "Zentrale/Standort", *namen[1:]]
    for pflicht in ("Kanzleiname", "Rechtsform", "Straße", "PLZ", "Ort", "Bundesland", "Telefon",
                    "E-Mail", "Website", "Hauptansprechpartner", "Teamgröße", "Teamgröße-Basis",
                    "Konfidenz", "Quell-URLs", "Abrufdatum", "Status"):
        assert pflicht in kopf
    assert ws.max_row == 21


def test_kanzleien_filter_und_kopfzeile_fixiert(wb):
    ws = wb[BLATT_KANZLEIEN]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref.startswith("A1:")
    assert ws.auto_filter.ref.endswith("21")
    assert ws["A1"].font.bold


def test_status_dropdown(wb):
    ws = wb[BLATT_KANZLEIEN]
    status_spalte = _kopf(ws).index("Status") + 1
    buchstabe = ws.cell(row=1, column=status_spalte).column_letter
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    dv = dvs[0]
    assert dv.type == "list"
    assert dv.formula1 == '"neu,geprüft,kontaktiert,kein Interesse,ungültig"'
    assert str(dv.sqref) == f"{buchstabe}2:{buchstabe}{DROPDOWN_MIN_ZEILEN}"


def test_kanzlei_werte(wb):
    zeilen = {z["Domain"]: z for z in _zeilen(wb[BLATT_KANZLEIEN])}
    z = zeilen["hansen-steuerberatung.example"]
    assert z["Kanzleiname"] == "Hansen & Petersen Steuerberater PartG mbB"
    assert z["PLZ"] == "24103"  # als Text
    assert z["Telefon"] == "+49 431 9987610"
    assert z["Hauptansprechpartner"] == "Jens Hansen"
    assert z["Titel"] == "Dipl.-Kfm."
    assert z["Funktion"] == "Steuerberater, Partner"
    assert z["Weitere Ansprechpartner"] == (
        "Maren Petersen (Steuerberaterin, Partnerin); M.Sc. Ole Brandt (Steuerberater)"
    )
    assert z["Teamgröße"] == 18
    assert z["Größenklasse"] == "11-20"  # gespeichert „16-50" (alte Einteilung), aus der Zahl neu
    assert z["Konfidenz"] == "hoch"
    assert z["Quell-URLs"].splitlines() == [
        "https://www.hansen-steuerberatung.example/impressum",
        "https://www.hansen-steuerberatung.example/team",
    ]
    assert z["Abrufdatum"].date() == date(2026, 10, 20)
    assert z["Status"] == "neu"


def test_teamgroesse_klasse_wenn_keine_zahl(wb):
    zeilen = {z["Domain"]: z for z in _zeilen(wb[BLATT_KANZLEIEN])}
    # gespeichert mit alter Einteilung: Text-Basis neu ausgewertet („über 20" → 21-50, „rund 30" → 21-50)
    assert zeilen["foerde-steuer.example"]["Teamgröße"] == "21-50"
    assert zeilen["ostholstein-steuer.example"]["Größenklasse"] == "21-50"
    assert zeilen["nordtax-norderstedt.example"]["Größenklasse"] == "51-100"
    # nicht umrechenbar („gut zehn", „kleine Kanzlei"): leer, mit Hinweis
    for d in ("wedel-steuer.example", "dithmarscher-steuerkanzlei.example"):
        assert zeilen[d]["Größenklasse"] is None and "alte Einteilung" in zeilen[d]["Abweichung"]
    assert zeilen["steuerbuero-schuett.example"]["Teamgröße"] is None


def test_fehlende_felder_bleiben_leer(wb):
    zeilen = {z["Domain"]: z for z in _zeilen(wb[BLATT_KANZLEIEN])}
    z = zeilen["stb-oldesloe.example"]
    assert z["Kanzleiname"] is None
    assert z["Hauptansprechpartner"] is None
    assert z["Weitere Ansprechpartner"] is None
    assert z["Konfidenz"] == "niedrig"


def test_nicht_gefunden_blatt(wb):
    zeilen = _zeilen(wb[BLATT_NICHT_GEFUNDEN])
    assert _kopf(wb[BLATT_NICHT_GEFUNDEN])[:2] == ["Domain", "Grund"]
    assert {z["Domain"]: z["Grund"] for z in zeilen}["stb-ratzeburg.example"] == "robots.txt verbietet Zugriff"
    assert len(zeilen) == 3


def test_dubletten_blatt(wb):
    zeilen = _zeilen(wb[BLATT_DUBLETTEN])
    assert len(zeilen) == 2
    assert {(z["Behalten"], z["Dublette"], z["Kriterium"], z["Art"]) for z in zeilen} == {
        ("nordsee-steuerberatung.example", "nordsee-steuer-husum.example", "telefon", "zusammengeführt"),
        ("hansen-steuerberatung.example", "hansen-petersen-kiel.example", "plz+name", "unsicher"),
    }


def test_lauf_protokoll_blatt(wb):
    ws = wb[BLATT_PROTOKOLL]
    for spalte in ("Suchbegriff", "Trefferzahl", "Laufzeit (s)", "Anfragen", "Kosten", "Neue Kandidaten"):
        assert spalte in _kopf(ws)
    zeilen = _zeilen(ws)
    assert len(zeilen) == 4
    assert zeilen[0]["Suchbegriff"] == "Steuerberater Kiel"
    assert zeilen[0]["Trefferzahl"] == 20
    assert zeilen[0]["Laufzeit (s)"] == pytest.approx(2.84)
    assert zeilen[2]["Fehler"].startswith("brave:")
    assert zeilen[0]["Anfragen"] == 2
    assert zeilen[0]["Kosten"] == pytest.approx(0.01)


def test_alle_blaetter_mit_filter_und_fixierung(wb):
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref


def test_leerer_export(tmp_path):
    with Storage(tmp_path / "leer.sqlite") as st:
        pfad = export_excel(st, "SH", tmp_path / "leer.xlsx")
    wb = load_workbook(pfad)
    assert len(wb.sheetnames) == 4
    assert wb[BLATT_KANZLEIEN].max_row == 1
    assert len(wb[BLATT_KANZLEIEN].data_validations.dataValidation) == 1


def test_export_filtert_bundesland(befuellt, tmp_path):
    pfad = export_excel(befuellt, "HH", tmp_path / "hh.xlsx")
    wb = load_workbook(pfad)
    assert wb[BLATT_KANZLEIEN].max_row == 1
    assert wb[BLATT_PROTOKOLL].max_row == 1


def test_neue_spalten(befuellt, tmp_path):
    from recherche.models import Ansprechpartner, Kanzlei, NichtGefunden

    befuellt.upsert_kanzlei(Kanzlei(lauf_bundesland="SH", kandidat_id=99, kanzleiname="Kanzlei Ohneweb",
                                    website_vorhanden=False, abweichung="Telefon: places 1 / Website 2",
                                    ansprechpartner=[Ansprechpartner(name="A", funktion="StB")], berufstraeger_anzahl=1))
    befuellt.add_nicht_gefunden(NichtGefunden(domain="weg.example", bundesland="SH", grund="Timeout", name="Kanzlei Weg",
                                              strasse="Wall 5", plz="24103", ort="Kiel", telefon="0431 1"))
    wb = load_workbook(export_excel(befuellt, "SH", tmp_path / "x.xlsx"))
    zeilen = {z["Kanzleiname"]: z for z in _zeilen(wb[BLATT_KANZLEIEN])}
    z = zeilen["Kanzlei Ohneweb"]
    assert (z["Website vorhanden"], z["Berufsträger (Anzahl)"], z["Abweichung"]) == ("nein", 1, "Telefon: places 1 / Website 2")
    assert zeilen["Steuerbüro Schütt"]["Website vorhanden"] == "ja"
    ng = {r["Domain"]: r for r in _zeilen(wb[BLATT_NICHT_GEFUNDEN])}["weg.example"]
    assert (ng["Name (Suche)"], ng["Straße"], ng["PLZ"], ng["Ort"], ng["Telefon"]) == ("Kanzlei Weg", "Wall 5", "24103", "Kiel", "0431 1")
