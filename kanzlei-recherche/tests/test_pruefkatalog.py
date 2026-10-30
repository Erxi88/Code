"""Prüfkatalog (docs/pruefkatalog.md), Abschnitt 5.1: Negativ-Fixtures für 1A und die Feldregeln aus 2.

Jeder Fall in tests/fixtures/negativ/*.yaml ist fiktiv: Impressum-Zeilen, die LLM-Antwort (von Hand) und das
erwartete Ergebnis. Durchlaufen wird die echte Kette ohne Ollama und ohne Netz: Regeln → Zusammenführung →
Kanzlei-Prüfung mit Bundesland-Filter.
"""

from pathlib import Path

import pytest
import yaml

from recherche.extract.base import LlmErgebnis, LlmPerson
from recherche.extract.kanzlei import grund_aus_abweichung, pruefen
from recherche.extract.lauf import regeln_fuer
from recherche.extract.merge import setze_ortsliste, zusammenfuehren
from recherche.models import Kandidat, Seite, Seitentyp
from recherche.normalize import PlzPruefung
from recherche.search.pipeline import Ort

FAELLE = sorted((Path(__file__).parent / "fixtures" / "negativ").glob("*.yaml"))
SH = "Schleswig-Holstein"
PLZ = {"24103": "Kiel", "24114": "Kiel", "25813": "Husum", "24768": "Rendsburg", "21483": "Lütau"}
BLOCKLIST = {"steuer-portal-blockiert.example"}


def _pruefe(plz, ort):
    if plz in PLZ:
        return PlzPruefung(True, SH, orte=[PLZ[plz]], ort_kanonisch=PLZ[plz])
    return PlzPruefung(False, None, hinweis=f"PLZ {plz} unbekannt")


@pytest.fixture
def ausschluss(settings):
    from recherche.extract.rules import Ausschluss

    return Ausschluss.laden(settings.pfad(settings.pfade.impressum_ausschluss))


@pytest.fixture(autouse=True)
def ortsliste():
    setze_ortsliste([Ort("Rendsburg", "24768"), Ort("Kiel", "24103", plz_weitere=("24105", "24114"))])
    yield
    setze_ortsliste([])


def _html(zeilen):
    return "<html><body>" + "".join(f"<p>{z}</p>" for z in zeilen) + "</body></html>"


def _lauf(fall: dict, ausschluss):
    domain = fall["domain"]
    seiten = {Seitentyp.IMPRESSUM: Seite(url=f"https://{domain}/impressum", domain=domain, seitentyp=Seitentyp.IMPRESSUM,
                                         http_status=200, text="\n".join(fall["impressum"]), html=_html(fall["impressum"]))}
    llm = LlmErgebnis(**{**fall.get("llm", {}),
                         "personen": [LlmPerson(**p) for p in fall.get("llm", {}).get("personen", [])]})
    kandidat = Kandidat(id=1, bundesland="SH", domain=domain, website=f"https://{domain}/")
    k = zusammenfuehren(kandidat, [], seiten, regeln_fuer(seiten, ausschluss, domain), llm, _pruefe, ausschluss)
    return pruefen(k, seiten, funde=[], ziel_land=SH, blocklist=BLOCKLIST)


def test_mindestumfang_der_negativ_fixtures():
    """Abschnitt 5.1: mindestens 15 Negativ-Fixtures, je Ausschlusskategorie aus 1A eine."""
    kategorien = {yaml.safe_load(p.read_text(encoding="utf-8"))["kategorie"].split(" (")[0] for p in FAELLE}
    assert len(FAELLE) >= 15
    for k in ("1A Lohnsteuerhilfeverein", "1A Rechtsanwaltskanzlei", "1A Buchhaltungsbüro", "1A Institution",
              "1A Behörde", "1A Verein/Partei", "1A Portal", "1A Branchenfremd", "1A Ausland"):
        assert k in kategorien, k


@pytest.mark.parametrize("pfad", FAELLE, ids=[p.stem for p in FAELLE])
def test_negativ_fixture(pfad, ausschluss):
    fall = yaml.safe_load(pfad.read_text(encoding="utf-8"))
    soll = fall["erwartet"]
    k = _lauf(fall, ausschluss)
    assert k.status.value == soll["status"], k.abweichung
    if "grund" in soll:
        grund = grund_aus_abweichung(k.abweichung) if soll["status"] == "keine Kanzlei" else k.abweichung
        assert grund and soll["grund"] in grund, k.abweichung
    if "ansprechpartner" in soll:
        assert [p.name for p in k.ansprechpartner] == soll["ansprechpartner"]
    if "telefon" in soll:
        assert k.telefon == soll["telefon"]
    if "plz" in soll:
        assert k.plz == soll["plz"]
    for teil in soll.get("abweichung", []):
        assert teil in (k.abweichung or ""), k.abweichung


# ---------------------------------------------------------------------- Ansprechpartner (Abschnitt 2)


@pytest.mark.parametrize(("name", "titel", "funktion", "erwartet"), [
    ("Porträtfoto von Jan Hansen", None, "Steuerberater", ("Jan Hansen", None, "Steuerberater")),
    ("Foto von Eva Kittel", None, None, ("Eva Kittel", None, None)),
    ("Bild: Eva Kittel", None, None, ("Eva Kittel", None, None)),
    ("Jan Hansen Steuerberater", None, None, ("Jan Hansen", None, "Steuerberater")),
    ("Jan Hansen, StB", None, "Partner", ("Jan Hansen", None, "Partner, StB")),
    ("Jan Hansen Dipl.-Kfm.", None, None, ("Jan Hansen", "Dipl.-Kfm.", None)),
])  # fmt: skip
def test_ansprechpartner_praefixe_und_suffixe(name, titel, funktion, erwartet):
    from recherche.extract.ansprechpartner import bereinigen
    from recherche.models import Ansprechpartner

    (p,), hinweise = bereinigen([Ansprechpartner(name=name, titel=titel, funktion=funktion)])
    assert (p.name, p.titel, p.funktion) == erwartet and not hinweise


@pytest.mark.parametrize(("name", "funktion", "bleibt"), [
    ("Oliver Hagen", "Steuerberater", True),  # Ortsname als Nachname bleibt (Rücksprache 30.10.2026)
    ("August Petersen", None, True),  # Monatsname als Vorname bleibt
    ("Jan Lund Kiel", "Steuerberater", False),  # Ortsname als Zusatzwort
    ("Fachanwalt für Familienrecht Elmshorn", None, False),
    ("Ina Lund", "Amtsdirektorin", False),
    ("Archiv", None, False),
    ("Montag Dienstag", None, False),  # nur Wochentage
    ("Hanh Mai Luong", "Steuerberaterin", True),  # Monatsname ohne Zahl ist ein Name (Rücksprache 30.10.2026)
    ("Jan Lund 3. Mai", None, False),  # Monatsname mit Zahl ist ein Datum
    ("Jan Lund Mai 2024", None, False),
    ("Schröder", "Steuerberaterin", False),  # nur Nachname
    ("Dr. Schröder", None, False),  # erstes Wort ein Titel
    ("W. Drews", "Steuerberaterin", True),  # Initiale ist ein Vorname
])  # fmt: skip
def test_ansprechpartner_verwerfen(name, funktion, bleibt):
    from recherche.extract.ansprechpartner import bereinigen
    from recherche.models import Ansprechpartner

    liste, hinweise = bereinigen([Ansprechpartner(name=name, funktion=funktion)])
    assert bool(liste) is bleibt and bool(hinweise) is not bleibt


def test_rechtsanwaelte_nur_ohne_steuerberater_vorn():
    from recherche.extract.ansprechpartner import bereinigen
    from recherche.models import Ansprechpartner

    liste, _ = bereinigen([Ansprechpartner(name="Ole Prasse", funktion="Rechtsanwalt, Notar"),
                           Ansprechpartner(name="Ina Lund", funktion="Steuerberaterin")])
    assert [p.name for p in liste] == ["Ina Lund", "Ole Prasse"]
    liste, _ = bereinigen([Ansprechpartner(name="Ole Prasse", funktion="Rechtsanwalt")])
    assert [p.name for p in liste] == ["Ole Prasse"]  # ohne Steuerberater bleibt er vorn


# ---------------------------------------------------------------------- Telefon (Abschnitt 2)


def test_telefon_rangfolge_tel_link_schema_impressum_kontakt():
    from recherche.extract.rules import extract_impressum

    impressum = "Steuerkanzlei Lund\nMarkt 1\n24103 Kiel\nTel. 0431 1111111"
    kontakt = ("Kontakt\nTelefon 0431 2222222", [])
    # tel:-Link vor Schema.org vor Impressum-Text vor Kontaktseite – unter den Nummern, die auch sichtbar im Text stehen
    assert extract_impressum(impressum, links=["tel:+494312222222"], weitere_texte=[kontakt],
                             schema_telefone=["+49 431 1111111"]).telefon == "+49 431 2222222"
    assert extract_impressum(impressum, weitere_texte=[kontakt], schema_telefone=["+49 431 2222222"]).telefon == "+49 431 2222222"
    assert extract_impressum(impressum, weitere_texte=[kontakt]).telefon == "+49 431 1111111"
    # Link/Schema.org ohne sichtbare Nummer gilt nur, wenn der Text keine nennt
    assert extract_impressum("Markt 1\n24103 Kiel", links=["tel:+494313333333"]).telefon == "+49 431 3333333"
    assert extract_impressum("Markt 1\n24103 Kiel", schema_telefone=["+49 431 4444444"]).telefon == "+49 431 4444444"
    assert extract_impressum("Markt 1\n24103 Kiel", weitere_texte=[kontakt]).telefon == "+49 431 2222222"


def test_telefon_festnetz_vor_mobil_vor_service_und_kein_fax():
    from recherche.extract.rules import extract_impressum

    text = "Steuerkanzlei Lund\nMarkt 1\n24103 Kiel\nHotline 0800 1234567\nMobil: 0171 1234567\nF: 0431 7654321"
    erg = extract_impressum(text, links=["tel:08001234567", "tel:01711234567"])
    assert erg.telefon == "+49 171 1234567"  # kein Festnetz: Mobil vor 0800; „F:" ist Fax
    assert extract_impressum(text + "\nTel. 0431 5550000", links=["tel:01711234567"]).telefon == "+49 431 5550000"


def test_tel_link_aus_dem_kammer_abschnitt_zaehlt_nicht(ausschluss):
    from recherche.extract.rules import extract_impressum

    text = ("Steuerkanzlei Lund\nMarkt 1\n24103 Kiel\nTel. 0431 5550000\n"
            "Zuständige Kammer: Steuerberaterkammer Schleswig-Holstein\nHopfenstraße 2 d\n24114 Kiel\nTel. 0431 5703490")
    assert extract_impressum(text, ausschluss, links=["tel:+494315703490"]).telefon == "+49 431 5550000"


@pytest.mark.parametrize(("roh", "erwartet"), [
    ("0 41 53 / 5 97 40", "+49 4153 59740"),
    ("+49 (0) 41 53 59740", "+49 4153 59740"),
    ("0431 1", None),  # zu kurz
    ("0431 12", None),  # laut phonenumbers ungültig
])  # fmt: skip
def test_telefon_nur_gueltige_nummern(roh, erwartet):
    from recherche.normalize import normalize_telefon

    assert normalize_telefon(roh) == erwartet


# ---------------------------------------------------------------------- 1B und Qualitätsbericht (Abschnitt 4)


def _k(**kw):
    from recherche.models import Kanzlei

    return Kanzlei(**{"lauf_bundesland": "SH", "domain": "lund.example", "kanzleiname": "Lund Beratung", "plz": "24103",
                      "ort": "Kiel", "bundesland": SH, "plz_verifiziert": True, **kw})


def _impressum(*zeilen):
    return {Seitentyp.IMPRESSUM: Seite(url="https://lund.example/impressum", domain="lund.example",
                                       seitentyp=Seitentyp.IMPRESSUM, text="\n".join(zeilen), html=_html(zeilen))}


@pytest.mark.parametrize(("kw", "zeilen", "funde", "rang", "konfidenz"), [
    ({}, ["Lund Beratung", "Steuerberaterkammer Schleswig-Holstein"], [], 1, "hoch"),
    ({"ansprechpartner": [{"name": "Anke Lund", "funktion": "Steuerberaterin"}]}, ["Lund Beratung"], [], 2, "hoch"),
    ({}, ["Lund Beratung", "Berufsbezeichnung: Steuerberater (verliehen in Deutschland)"], [], 2, "hoch"),
    ({"kanzleiname": "Lund Steuerberatungsgesellschaft mbH"}, ["Lund"], [], 3, "hoch"),
    ({"kanzleiname": "Steuerkanzlei Lund"}, ["Steuerkanzlei Lund"], [], 4, "mittel"),
    ({}, ["Lund Beratung"], [{"quelle": "brave", "snippet": "Ihr Steuerberater in Kiel"}], 5, "niedrig"),
    ({"plz_verifiziert": None}, ["Lund Beratung", "Steuerberaterkammer Schleswig-Holstein"], [], 1, "mittel"),
    ({"abweichung": "Telefon ungültig: 0431 1 (Website)"}, ["Steuerberaterkammer"], [], 1, "mittel"),  # Q4-Deckel
])  # fmt: skip
def test_konfidenz_nach_1b(kw, zeilen, funde, rang, konfidenz):
    from recherche.extract.kanzlei import signal_1b
    from recherche.models import Ansprechpartner, Fund

    kw = {**kw, "ansprechpartner": [Ansprechpartner(**p) for p in kw.get("ansprechpartner", [])]}
    f = [Fund(bundesland="SH", quell_id=str(i), lauf_id="l", **x) for i, x in enumerate(funde)]
    k = pruefen(_k(**kw), _impressum(*zeilen), funde=f, ziel_land=SH)
    assert signal_1b(k, "\n".join(zeilen), f)[0] == rang and k.konfidenz.value == konfidenz, k.abweichung
    assert ("1B Rang" in (k.abweichung or "")) is (rang >= 4)


def test_ohne_website_aus_dem_amtlichen_verzeichnis_ist_hoch():
    from recherche.extract.kanzlei import anwenden
    from recherche.models import Fund

    f = [Fund(bundesland="SH", quelle="stb_verzeichnis", quell_id="x", lauf_id="l", snippet="Amtliches …")]
    k = anwenden(_k(domain=None, website_vorhanden=False), None, False, SH, funde=f)
    assert k.konfidenz.value == "hoch" and k.status.value == "neu"


def test_zu_pruefen_q5_q6_q12_und_1b():
    from recherche.extract.ansprechpartner import setze_ortsnamen
    from recherche.models import Konfidenz, Status
    from recherche.qualitaet import zu_pruefen_grund

    setze_ortsnamen(["Kiel"])
    hoch = Konfidenz.HOCH
    assert zu_pruefen_grund(_k(konfidenz=hoch)) is None
    assert zu_pruefen_grund(_k(konfidenz=Konfidenz.NIEDRIG)) == "Q12 Konfidenz niedrig"
    assert zu_pruefen_grund(_k(konfidenz=Konfidenz.NIEDRIG, status=Status.GEPRUEFT)) is None  # gepflegt bleibt
    assert zu_pruefen_grund(_k(konfidenz=hoch, kanzleiname="Kontakt")).startswith("Q6 Navigationswort")
    assert zu_pruefen_grund(_k(konfidenz=hoch, kanzleiname="Kiel")) == "Q6 nur Ortsname „Kiel“"
    assert zu_pruefen_grund(_k(konfidenz=hoch, kanzleiname="Die Amtsdirektorin")).startswith("Q6 Rolle")
    rang4 = "1B Rang 4: nur Kanzleiwort im Namen (keine Kammer, keine Berufsbezeichnung)"
    assert zu_pruefen_grund(_k(konfidenz=Konfidenz.MITTEL, abweichung=rang4)).startswith("1B Rang 4")
    abgeleitet = f"{rang4}; Name abgeleitet aus Ansprechpartner Ina Lund – kein Kanzleiname gefunden"
    assert zu_pruefen_grund(_k(konfidenz=Konfidenz.MITTEL, abweichung=abgeleitet)).startswith("Q5 Name abgeleitet")


def test_qualitaetsbericht_alle_pruefungen():
    from recherche.models import Ansprechpartner, Dublette, Konfidenz, Status
    from recherche.qualitaet import bericht

    hoch = Konfidenz.HOCH
    stb = [Ansprechpartner(name="Maik Lund", funktion="Steuerberater")]
    zeilen = [
        _k(domain="at.example", status=Status.AUSSERHALB, abweichung="Außerhalb: Ausland (.at)"),
        _k(domain="lsthv.example", status=Status.KEINE_KANZLEI, abweichung="Keine Kanzlei: Lohnsteuerhilfeverein"),
        _k(domain="a.example", konfidenz=hoch, abweichung="Ansprechpartner ohne Vorname: Schröder; Telefon ungültig: 0431 1",
           email="datenschutz@a.example", strasse="Markt 1"),
        _k(domain="b.example", konfidenz=hoch, kanzleiname="Steuerbüro B", strasse="Markt 1", abweichung="PLZ/Ort: Kiel ≠ Kie"),
        _k(domain="c.example", konfidenz=Konfidenz.NIEDRIG),
        _k(domain="d.example", konfidenz=hoch, teamgroesse=250, abweichung="Teamgröße unplausibel (250 Personen, 2 Berufsträger)"),
        _k(domain="zentrale.example", kanzleiname="Nordbund StBG", konfidenz=hoch, telefon="+49 431 59360", ansprechpartner=stb),
        _k(domain="heide.zentrale.example", kanzleiname="Nordbund Heide", konfidenz=hoch, telefon="+49 431 59360",
           ansprechpartner=stb, strasse="Wall 2", plz="25746"),
    ]
    dub = [Dublette(behalten="x.example", dublette="y.example", bundesland="SH", kriterium="telefon+ansprechpartner")]
    erg = bericht(zeilen, dub, SH, coverage_gefunden=9)
    n = erg.anzahl()
    assert {nr: n.get(nr, 0) for nr in ("Q1", "Q2", "Q3", "Q4", "Q7", "Q8", "Q9", "Q10", "Q11", "Q12", "Q13", "Q14")} == {
        "Q1": 1, "Q2": 1, "Q3": 1, "Q4": 1, "Q7": 1, "Q8": 1, "Q9": 1, "Q10": 1, "Q11": 1, "Q12": 1, "Q13": 1, "Q14": 1}
    assert list(erg.zu_pruefen.values()) == ["Q12 Konfidenz niedrig"]
    q14 = next(b.text for b in erg.befunde if b.nr == "Q14")
    assert q14 == "coverage 9, Hauptblatt 5 (Differenz 4)"


def test_qualitaet_befehl(tmp_path, db_pfad):
    from typer.testing import CliRunner

    from recherche import cli
    from recherche.models import Konfidenz
    from recherche.storage import Storage

    with Storage(db_pfad) as st:
        st.upsert_kanzlei(_k(konfidenz=Konfidenz.NIEDRIG))
        st.upsert_kanzlei(_k(domain="b.example", konfidenz=Konfidenz.HOCH, abweichung="Telefon ungültig: 1"))
    out = tmp_path / "q.csv"
    r = CliRunner().invoke(cli.app, ["qualitaet", "-b", "SH", "--db", str(db_pfad), "--out", str(out)])
    assert r.exit_code == 0, r.output
    assert "Qualitätsbericht SH" in r.output and "Zu prüfen (aus dem Hauptblatt verschoben): 1" in r.output
    assert "Q12;Konfidenz niedrig;Sperre" in out.read_text(encoding="utf-8-sig")


# ---------------------------------------------------------------------- Q4 Vorwahl ↔ PLZ (BNetzA)

VORWAHLEN = Path(__file__).parent / "fixtures" / "vorwahlen"


@pytest.fixture
def vorwahl_pruefer():
    from recherche.vorwahl import VorwahlPruefer

    orte = [Ort("Kiel", "24103", "Kreisfreie Stadt Kiel", plz_weitere=("24105",)),
            Ort("Kronshagen", "24119", "Kreis Rendsburg-Eckernförde"),
            Ort("Pinneberg", "25421", "Kreis Pinneberg"), Ort("Quickborn (Pinneberg)", "25451", "Kreis Pinneberg"),
            Ort("Quickborn (Dithmarschen)", "25712", "Kreis Dithmarschen"), Ort("Husum", "25813", "Kreis Nordfriesland"),
            Ort("Fehmarn", "23769", "Kreis Ostholstein"), Ort("Burg (Dithmarschen)", "25712", "Kreis Dithmarschen"),
            Ort("Sylt", "25980", "Kreis Nordfriesland")]
    return VorwahlPruefer.laden(VORWAHLEN / "onb_ausschnitt.csv", orte, VORWAHLEN / "nachbarkreise.yaml", "SH")


@pytest.mark.parametrize(("telefon", "plz", "passt"), [
    ("+49 431 1234567", "24103", True),  # gleicher Kreis
    ("+49 431 1234567", "24119", True),  # Kronshagen: Nachbarkreis von Kiel (Rücksprache 30.10.2026)
    ("+49 40 1234567", "25421", True),  # Hamburg 040 für SH zulässig
    ("+49 4106 123456", "25451", True),  # „Quickborn Kr Pinneberg" → Quickborn im Kreis Pinneberg
    ("+49 2357 123456", "24103", False),  # Herscheid (NRW) bei Kiel: Vorwahlzone passt nicht
    ("+49 4841 123456", "24103", False),  # Husum bei Kiel: Kreis Nordfriesland ist kein Nachbar von Kiel
    ("+49 4371 123456", "23769", True),  # „Burg auf Fehmarn" ist nicht Burg (Dithmarschen) – nicht entscheidbar
    ("+49 4651 123456", "25980", True),  # Westerland ↔ Sylt: kein Gemeinde-Treffer, Zone 4 passt
    ("+49 171 1234567", "24103", True),  # Mobil wird nicht geprüft
])  # fmt: skip
def test_vorwahl_passt_zur_plz(vorwahl_pruefer, telefon, plz, passt):
    assert (vorwahl_pruefer.pruefen(telefon, plz) is None) is passt


def test_vorwahl_hinweis_in_der_kanzlei_pruefung(vorwahl_pruefer):
    from recherche.qualitaet import bericht
    from recherche.vorwahl import setze_pruefer

    setze_pruefer(vorwahl_pruefer)
    seiten = _impressum("Lund Beratung", "Steuerberaterkammer Schleswig-Holstein")
    k = pruefen(_k(telefon="+49 2357 123456"), seiten, ziel_land=SH)
    assert "Vorwahl passt nicht zur PLZ: 02357 (Herscheid Westf)" in k.abweichung and k.konfidenz.value == "mittel"
    k = pruefen(k, seiten, ziel_land=SH)
    assert k.abweichung.count("Vorwahl passt nicht") == 1  # wiederholte Prüfung setzt den Hinweis nicht doppelt
    assert [b.nr for b in bericht([k]).befunde] == ["Q4"]
    k.telefon = "+49 431 1234567"
    k = pruefen(k, seiten, ziel_land=SH)
    assert "Vorwahl" not in (k.abweichung or "") and k.konfidenz.value == "hoch"


# ---------------------------------------------------------------------- Nachbarkreise aus Wikidata + GV-ISys


def test_nachbarkreise_aus_wikidata_und_gv(tmp_path):
    import yaml as _yaml
    from typer.testing import CliRunner

    from recherche import cli

    gv = Path(__file__).parent / "fixtures" / "gv_isys" / "ausschnitt_sh.csv"
    xml = Path(__file__).parent / "fixtures" / "nachbarkreise" / "wikidata_ausschnitt.xml"
    kopie = tmp_path / "wikidata.xml"
    kopie.write_bytes(xml.read_bytes())
    ziel = tmp_path / "nachbarkreise.yaml"
    ziel.write_text(_yaml.safe_dump({"SH": {"immer_zulaessig": ["040"], "nachbarn": {
        "Kreis Dithmarschen": ["Kreis Nordfriesland", "Kreis Rendsburg-Eckernförde"],
        "Kreis Nordfriesland": ["Kreis Rendsburg-Eckernförde"]}}}, allow_unicode=True), encoding="utf-8")
    r = CliRunner().invoke(cli.app, ["nachbarkreise", "--wikidata", str(kopie), "--gv", str(gv), "--out", str(ziel),
                                     "--xml-loeschen"])
    assert r.exit_code == 0, r.output
    ausgabe = " ".join(r.output.split())  # rich bricht lange Zeilen um
    daten = _yaml.safe_load(ziel.read_text(encoding="utf-8"))
    sh = daten["SH"]
    assert sh["immer_zulaessig"] == ["040"]  # Ausnahme aus der alten Datei bleibt
    assert sh["nachbarn"]["Kreis Dithmarschen"] == ["Kreis Nordfriesland", "Kreis Rendsburg-Eckernförde",
                                                     "Kreis Schleswig-Flensburg"]
    assert "Kreis Dithmarschen" in sh["nachbarn"]["Kreis Schleswig-Flensburg"]  # symmetrisch
    assert "Kreisfreie Stadt Kiel" in sh["nachbarn"]["Kreis Rendsburg-Eckernförde"]  # Namen aus GV-ISys
    assert not any("Husum" in k or "Husum" in str(v) for k, v in sh["nachbarn"].items())  # Altkreis verworfen
    assert "1 Schlüssel nicht (mehr) im GV-ISys verworfen" in ausgabe
    assert "SH: neu aus Wikidata" in ausgabe and "Kreis Dithmarschen – Kreis Schleswig-Flensburg" in ausgabe
    assert "SH: nur in der bisherigen Liste" in ausgabe and "Kreis Nordfriesland – Kreis Rendsburg-Eckernförde" in ausgabe
    assert "SH: ohne Nachbarpaar" in ausgabe and "Kreisfreie Stadt Lübeck" in ausgabe
    assert "HH: ohne Nachbarpaar" in ausgabe and "Hamburg" in ausgabe
    assert not kopie.exists() and (tmp_path / "nachbarkreise.yaml.bak").exists()


# ---------------------------------------------------------------------- Rückmeldung sh_gesamt_v2 (30.10.2026)


@pytest.mark.parametrize(("name", "impressum", "personen", "grund"), [
    ("Steuerberaterverband Nordland e.V.", ["Impressum", "Steuerberaterverband Nordland e.V.", "Steuerberaterkammer Nordland"],
     [], "Institution"),  # nicht „Lohnsteuerhilfeverein" – Institution geht vor
    ("Steuerberater Gemeinde Musterbek", ["Impressum", "Gemeinde Musterbek", "Der Bürgermeister", "Dorfstraße 1"], [], "Behörde"),
    ("Steuerberater Ole Lund", ["Impressum", "Handels- und Gewerbeverein Musterstadt e.V.", "Markt 1"], [], "Verein/Partei"),
    ("Lohnsteuerhilfe Nord e. V.", ["Impressum", "Lohnsteuerhilfe Nord e. V."], [], "Lohnsteuerhilfeverein"),
    # Rang 2/3 verhindert den schwachen Verein-Ausschluss über „VR …" im Impressum (franz-sahm.de, taxlegis.de)
    ("Sahm Steuerberatungsgesellschaft mbH", ["Impressum", "Sahm Steuerberatungsgesellschaft mbH", "Geschäftsführer: Franz Sahm",
     "Mitglied im Steuerberaterverband e.V., VR 1234"], [{"name": "Franz Sahm", "funktion": "Steuerberater"}], None),
    ("Steuerkanzlei Passau", ["Impressum", "Steuerkanzlei Passau", "Zuständige Kammer: Steuerberaterkammer Niedersachsen",
     "Partner im Netzwerk e.V. (Vereinsregister Hannover)"], [{"name": "Jörg Passau", "funktion": "Steuerberater"}], None),
])  # fmt: skip
def test_ausschluss_reihenfolge_und_rang(name, impressum, personen, grund):
    from recherche.extract.kanzlei import ausschluss_grund
    from recherche.models import Ansprechpartner

    k = _k(kanzleiname=name, ansprechpartner=[Ansprechpartner(**p) for p in personen])
    ergebnis = ausschluss_grund(k, "\n".join(impressum))
    assert (ergebnis.split(" (")[0] if ergebnis else None) == grund


# ---------------------------------------------------------------------- Stichprobe (Abschnitt 5)


def _hauptzeilen(n):
    from recherche.models import Konfidenz

    return [_k(id=i, domain=f"k{i:02}.example", kanzleiname=f"Steuerkanzlei {i}", konfidenz=Konfidenz.HOCH,
               telefon="+49 431 1234567", quell_urls=[f"https://k{i:02}.example/impressum"]) for i in range(n)]


def test_stichprobe_reproduzierbar_nur_hauptblatt():
    from recherche.models import Konfidenz, Status
    from recherche.stichprobe import ziehen

    zeilen = _hauptzeilen(60) + [_k(id=99, domain="zu.example", konfidenz=Konfidenz.NIEDRIG),
                                 _k(id=98, domain="weg.example", konfidenz=Konfidenz.HOCH, status=Status.KEINE_KANZLEI)]
    a, b = ziehen(zeilen, 50, seed=1), ziehen(list(reversed(zeilen)), 50, seed=1)
    assert len(a) == 50 and [k.domain for k in a] == [k.domain for k in b]  # Reihenfolge der Eingabe egal
    assert not {"zu.example", "weg.example"} & {k.domain for k in a}  # Zu prüfen / keine Kanzlei nie
    assert [k.domain for k in ziehen(zeilen, 50, seed=2)] != [k.domain for k in a]


def _ausfuellen(pfad, falsch_name=0, kanzlei_nein=0, offen=False):
    from openpyxl import load_workbook

    wb = load_workbook(pfad)
    ws = wb["Stichprobe"]
    kopf = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        for spalte in (c for c in kopf if c.endswith("– Prüfung")):
            ws.cell(row=r, column=kopf.index(spalte) + 1, value="richtig")
        ws.cell(row=r, column=kopf.index("Kanzlei?") + 1, value="ja")
        ws.cell(row=r, column=kopf.index("Ausland?") + 1, value="nein")
    for r in range(2, 2 + falsch_name):
        ws.cell(row=r, column=kopf.index("Name – Prüfung") + 1, value="falsch")
    for r in range(2, 2 + kanzlei_nein):
        ws.cell(row=r, column=kopf.index("Kanzlei?") + 1, value="nein")
    if offen:
        ws.cell(row=2, column=kopf.index("Telefon – Prüfung") + 1).value = None  # cell(value=None) ändert nichts
    wb.save(pfad)


@pytest.mark.parametrize(("falsch", "nein", "offen", "bestanden"), [
    (2, 0, False, True), (3, 0, False, False), (0, 1, False, False), (0, 0, True, False),
])  # fmt: skip
def test_stichprobe_abnahmekriterium(tmp_path, falsch, nein, offen, bestanden):
    from recherche.stichprobe import auswerten, schreiben, ziehen

    pfad = schreiben(ziehen(_hauptzeilen(60), 50, 1), tmp_path / "s.xlsx", "SH", 1)
    _ausfuellen(pfad, falsch, nein, offen)
    erg = auswerten(pfad)
    assert (erg.zeilen, erg.grenze, erg.falsch["Name"], erg.bestanden) == (50, 2, falsch, bestanden)


def test_stichprobe_befehl(tmp_path, db_pfad):
    from openpyxl import load_workbook
    from typer.testing import CliRunner

    from recherche import cli
    from recherche.storage import Storage

    with Storage(db_pfad) as st:
        for k in _hauptzeilen(12):
            st.upsert_kanzlei(k.model_copy(update={"id": None}))
    out = tmp_path / "stichprobe.xlsx"
    r = CliRunner().invoke(cli.app, ["stichprobe", "-b", "SH", "--db", str(db_pfad), "-n", "10", "--seed", "3", "--out", str(out)])
    assert r.exit_code == 0, r.output
    ws = load_workbook(out)["Stichprobe"]
    kopf = [c.value for c in ws[1]]
    assert ws.max_row == 11 and "Quell-URLs" in kopf and "Telefon – Prüfung" in kopf
    assert len(ws.data_validations.dataValidation) == 7  # 5 Kernfelder + Kanzlei? + Ausland?
    _ausfuellen(out)
    r = CliRunner().invoke(cli.app, ["stichprobe", "--auswerten", str(out)])
    assert r.exit_code == 0 and "Abnahme bestanden" in r.output
    _ausfuellen(out, falsch_name=1)  # 10 Zeilen: erlaubt 0 falsch
    r = CliRunner().invoke(cli.app, ["stichprobe", "--auswerten", str(out)])
    assert r.exit_code == 1 and "nicht bestanden" in r.output


# ---------------------------------------------------------------------- Ketten und Großgesellschaften (Abschnitt 9)


@pytest.fixture
def ketten_cfg():
    from recherche.ketten import KettenConfig, setze_config

    cfg = KettenConfig.laden(Path(__file__).parents[1] / "config" / "ketten.yaml")
    setze_config(cfg)
    yield cfg
    setze_config(KettenConfig())


def test_kanzleityp_und_ketten(ketten_cfg):
    from recherche.ketten import erkannte_ketten, kanzleityp

    zeilen = [
        _k(id=1, domain="pwc.de", kanzleiname="PricewaterhouseCoopers GmbH WPG"),
        _k(id=2, domain="mustermann-stb.example", kanzleiname="Steuerberater Ole Mustermann", rechtsform="Einzelkanzlei"),
        _k(id=3, domain="nord-partner.example", kanzleiname="Nord Partner mbB", rechtsform="PartG mbB"),
        _k(id=4, domain="heide.shbb.de", kanzleiname="SHBB Beratungsstelle Heide"),  # Kette laut ketten.yaml
        *[_k(id=10 + i, domain=f"steuerbuero-{o}.example", kanzleiname=f"Steuerbüro Kette {o}",
             quell_urls=["https://www.kette-zentrale.example/impressum"]) for i, o in enumerate(("nord", "sued", "ost"))],
        _k(id=20, domain="einzeln-a.example", quell_urls=["https://www.zwei.example/impressum"]),
        _k(id=21, domain="einzeln-b.example", quell_urls=["https://www.zwei.example/impressum"]),
    ]
    ketten = erkannte_ketten(zeilen, ketten_cfg)
    typ = {k.domain: kanzleityp(k, ketten, ketten_cfg) for k in zeilen}
    assert typ == {"pwc.de": "Großgesellschaft", "mustermann-stb.example": "Einzelkanzlei",
                   "nord-partner.example": "Mittelstand", "heide.shbb.de": "Kette",
                   "steuerbuero-nord.example": "Kette", "steuerbuero-sued.example": "Kette",
                   "steuerbuero-ost.example": "Kette",  # gleicher Impressum-Anbieter, 3 Domains
                   "einzeln-a.example": "Mittelstand", "einzeln-b.example": "Mittelstand"}  # nur 2 → keine Kette


def test_grossgesellschaft_ist_kein_ausschluss(ketten_cfg):
    from recherche.models import Ansprechpartner
    from recherche.search.pipeline import load_blocklist

    blocklist = load_blocklist(Path(__file__).parents[1] / "config" / "domain_blocklist.txt")
    assert not {"pwc.de", "deloitte.com", "bdo.de", "bakertilly.de"} & blocklist
    k = pruefen(_k(domain="pwc.de", kanzleiname="PricewaterhouseCoopers GmbH WPG",
                   ansprechpartner=[Ansprechpartner(name="Anke Lund", funktion="Steuerberaterin, Partnerin")]),
                _impressum("PricewaterhouseCoopers GmbH Wirtschaftsprüfungsgesellschaft", "Steuerberaterkammer Berlin"),
                ziel_land=SH, blocklist=blocklist)
    assert k.status.value == "neu" and k.konfidenz.value == "hoch"


def test_ketten_im_qualitaetsbericht_und_export(ketten_cfg, tmp_path):
    from openpyxl import load_workbook

    from recherche.export import BLATT_KANZLEIEN, build_workbook
    from recherche.models import Konfidenz
    from recherche.qualitaet import bericht

    zeilen = [_k(id=10 + i, domain=f"steuerbuero-{o}.example", kanzleiname=f"Steuerbüro Kette {o}", konfidenz=Konfidenz.HOCH,
                 quell_urls=["https://www.kette-zentrale.example/impressum"]) for i, o in enumerate(("nord", "sued", "ost"))]
    k1 = [b for b in bericht(zeilen).befunde if b.nr == "K1"]
    assert len(k1) == 1 and "3 Zeilen" in k1[0].text and k1[0].grad == "Hinweis"
    wb = build_workbook(zeilen, [], [], [], "Schleswig-Holstein")
    wb.save(tmp_path / "x.xlsx")
    ws = load_workbook(tmp_path / "x.xlsx")[BLATT_KANZLEIEN]
    kopf = [c.value for c in ws[1]]
    assert {r[kopf.index("Kanzleityp")] for r in ws.iter_rows(min_row=2, values_only=True)} == {"Kette"}


def test_regel_hinweise_fuer_telefon_und_ansprechpartner_rang():
    """fixture-eval soll zeigen, welche Regel ein Ergebnis bestimmt hat (Rückmeldung 30.10.2026)."""
    from recherche.extract.ansprechpartner import bereinigen
    from recherche.extract.rules import extract_impressum
    from recherche.models import Ansprechpartner

    # tel:-Link und sichtbarer Text verschieden: der Text gewinnt (Entscheidung 30.10.2026, sievertundpartner)
    erg = extract_impressum("Steuerkanzlei Lund\nMarkt 1\n24103 Kiel\nTel. 0431 980750", links=["tel:+49431960750"])
    assert erg.telefon == "+49 431 980750"
    assert "Telefon aus dem Text +49 431 980750, tel:-Link nennt +49 431 960750" in erg.hinweise
    gleich = extract_impressum("Markt 1\n24103 Kiel\nTel. 0431 980750", links=["tel:+49431980750"])
    assert not [h for h in gleich.hinweise if h.startswith("Telefon aus")]
    # Klasse geht weiter vor: Festnetz-Link schlägt eine Mobilnummer im Text
    mobil = extract_impressum("Markt 1\n24103 Kiel\nMobil 0171 1234567", links=["tel:+49431980750"])
    assert mobil.telefon == "+49 431 980750"
    _, hinweise = bereinigen([Ansprechpartner(name="Rolf Lund", funktion="Rechtsanwalt, Geschäftsführer"),
                              Ansprechpartner(name="Eva Kittel", funktion="Steuerberaterin")])
    assert hinweise == ["Ansprechpartner-Rang: Rolf Lund (Rechtsanwalt, Geschäftsführer) (Rechtsanwalt/Notar) hinter "
                        "Eva Kittel (Steuerberaterin) (Steuerberater)"]


def test_standortseite_einer_kette(tmp_path, ausschluss):
    """SHBB-Aufbau (fiktiv): Standort-Subdomain mit Kontaktblock und „Kanzleileitung", Impressum der Zentrale.
    Ansprechpartner = Standortleitung, Kanzleiname = Anbieter, Standortbezeichnung in „Zentrale/Standort"."""
    from recherche.extract.standorte import standort_spalte, standortbezeichnung
    from recherche.ketten import KettenConfig, setze_config

    (tmp_path / "ketten.yaml").write_text(
        "ketten:\n  nordbund:\n    name: Nordbund Steuerberatungsgesellschaft mbH\n    zentrale: nordbund.example\n"
        "    standortmuster: subdomain\n", encoding="utf-8")
    setze_config(KettenConfig.laden(tmp_path / "ketten.yaml"))
    try:
        start = ["Nordbund Steuerberatungskanzlei Heide", "Kontakt", "Nordbund Beratungsstelle Heide", "Markt 1",
                 "25746 Heide", "Telefon 0481 123450", "E-Mail info@heide.nordbund.example",
                 "Kanzleileitung und weitere Steuerberater", "Ole Lund", "Leiter", "Steuerberater",
                 "Durchwahl 0481 123451", "Eva Kittel", "Steuerberaterin", "Durchwahl 0481 123452"]
        impressum = ["Impressum", "Nordbund Steuerberatungsgesellschaft mbH", "Holstenstraße 1", "24103 Kiel",
                     "Telefon 0431 999990", "Geschäftsführer: Maik Muster",
                     "Zuständige Kammer: Steuerberaterkammer Schleswig-Holstein"]
        seiten = {
            Seitentyp.START: Seite(url="https://heide.nordbund.example/", domain="heide.nordbund.example",
                                   seitentyp=Seitentyp.START, text="\n".join(start), html=_html(start)),
            Seitentyp.IMPRESSUM: Seite(url="https://www.nordbund.example/impressum", domain="nordbund.example",
                                       seitentyp=Seitentyp.IMPRESSUM, text="\n".join(impressum), html=_html(impressum)),
        }
        llm = LlmErgebnis(kanzleiname="Nordbund Steuerberatungskanzlei Heide", verantwortliche_person="Maik Muster",
                          personen=[LlmPerson(name="Maik Muster", funktion="Geschäftsführer"),
                                    LlmPerson(name="Ole Lund", funktion="Leiter, Steuerberater")])
        kandidat = Kandidat(id=7, bundesland="SH", domain="heide.nordbund.example", website="https://heide.nordbund.example/")
        k = zusammenfuehren(kandidat, [], seiten, regeln_fuer(seiten, ausschluss, kandidat.domain), llm,
                            lambda p, o: PlzPruefung(True, SH, ort_kanonisch=o), ausschluss)
        assert k.hauptansprechpartner.name == "Ole Lund"
        assert (k.strasse, k.plz, k.telefon, k.email) == ("Markt 1", "25746", "+49 481 123450", "info@heide.nordbund.example")
        assert k.kanzleiname == "Nordbund Steuerberatungsgesellschaft mbH"
        assert standortbezeichnung(k) == "Nordbund Steuerberatungskanzlei Heide"
        zentrale = _k(id=8, domain="nordbund.example", kanzleiname="Nordbund Steuerberatungsgesellschaft mbH")
        spalte = standort_spalte([zentrale, k])
        assert spalte["heide.nordbund.example"] == ("Standort von Nordbund Steuerberatungsgesellschaft mbH "
                                                     "(Nordbund Steuerberatungskanzlei Heide)")
    finally:
        setze_config(KettenConfig())


def test_titel_vor_dem_namen_vor_der_vorname_pruefung():
    """900grad.de: „Dr. Dr. Rolf Harrie" ist vertretungsberechtigt und darf nicht als „ohne Vorname" fallen."""
    from recherche.extract.ansprechpartner import bereinigen
    from recherche.models import Ansprechpartner

    liste, hinweise = bereinigen([Ansprechpartner(name="Dr. Dr. Rolf Harrie", funktion="Geschäftsführer"),
                                  Ansprechpartner(name="Prof. Dr. rer. pol. Jan Hoff"),
                                  Ansprechpartner(name="WP/StB Dipl.-Kfm. Ina Kock"),
                                  Ansprechpartner(name="M.Sc. Tim Ott"), Ansprechpartner(name="B.A. Lea Ruge")])
    assert [(p.name, p.titel) for p in liste] == [("Rolf Harrie", "Dr. Dr."), ("Jan Hoff", "Prof. Dr. rer. pol."),
                                                   ("Ina Kock", "WP/StB Dipl.-Kfm."), ("Tim Ott", "M.Sc."),
                                                   ("Lea Ruge", "B.A.")]
    assert hinweise == []
    # nur Titel/Berufsbezeichnung + Nachname bleibt „ohne Vorname"
    for name in ("Dr. Schröder", "Steuerberaterin Schröder", "Herr Lund"):
        assert bereinigen([Ansprechpartner(name=name)])[0] == []


def test_fixture_eval_laedt_ketten_yaml(tmp_path):
    """fixture-eval/-prompt wenden die Standortseiten-Regel wie der Lauf an: ketten.yaml wird geladen – auch wenn
    meta.json keine Impressum-URL der Zentrale nennt (Standard-URL auf der Standort-Subdomain)."""
    from recherche.cli import _fixture_regeln_laden
    from recherche.config import Settings
    from recherche.ketten import standortseite

    (tmp_path / "ketten.yaml").write_text(
        "ketten:\n  nordbund:\n    name: Nordbund Steuerberatungsgesellschaft mbH\n    zentrale: nordbund.example\n",
        encoding="utf-8")
    s = Settings()
    s.pfade.ketten = str(tmp_path / "ketten.yaml")
    assert not standortseite("heide.nordbund.example", None, "https://heide.nordbund.example/impressum")
    _fixture_regeln_laden(s)
    assert standortseite("heide.nordbund.example", None, "https://heide.nordbund.example/impressum")
