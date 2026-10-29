import json

import pytest
from conftest import FIXTURES
from fakes import FakeOrtsAdapter, FakeQueryAdapter, hit
from openpyxl import load_workbook
from rich.console import Console
from typer.testing import CliRunner

from recherche import cli
from recherche.config import load_settings
from recherche.models import Fund, Kandidat
from recherche.search.pipeline import load_orte
from recherche.storage import Storage

runner = CliRunner()


@pytest.fixture(autouse=True)
def _breite_konsole(monkeypatch):
    monkeypatch.setattr(cli, "console", Console(width=300))


@pytest.fixture
def fakes(monkeypatch):
    brave = FakeQueryAdapter("brave", {
        "Steuerberater Kiel": ["https://a.example/", "https://www.facebook.com/x"],
        "Steuerberater Husum": ["https://h.example/"],
    }, kosten_je_1000=5.0)  # fmt: skip
    ddgs = FakeQueryAdapter("ddgs", {"Steuerberater Kiel": ["https://a.example/", "https://g.example/"]})
    places = FakeOrtsAdapter("places", {"Kiel": [hit("places", None, quell_id="p1", name="Kanzlei Ohneweb",
                                                     strasse="Wall 5", plz="24103", ort="Kiel")]})
    google = FakeQueryAdapter("google_cse", {}, grund="GOOGLE_CSE_API_KEY fehlt in .env", aktiv=False)
    alle = {"brave": brave, "ddgs": ddgs, "places": places, "google_cse": google}
    monkeypatch.setattr(cli, "build_adapter", lambda name, s, cache=None: alle[name])
    monkeypatch.setattr(cli, "aktive_adapter", lambda s, cache=None, stufe=None: [
        a for a in (brave, ddgs, places, google) if stufe is None or a.stufe == stufe])
    return alle


def _args(befehl, db_pfad, *rest):
    return [befehl, "-b", "SH", "--db", str(db_pfad), *rest]


def test_hilfe_listet_alle_befehle():
    r = runner.invoke(cli.app, ["--help"])
    assert r.exit_code == 0
    for befehl in ("search", "fill", "coverage", "crawl", "extract", "export", "run", "search-compare", "orte"):
        assert befehl in r.output


def test_export(befuellt, db_pfad, tmp_path):
    befuellt.close()
    out = tmp_path / "x.xlsx"
    r = runner.invoke(cli.app, _args("export", db_pfad, "--out", str(out)))
    assert r.exit_code == 0, r.output
    assert "20 Kanzleien" in r.output
    assert load_workbook(out)["Kanzleien"].max_row == 21


def test_unbekanntes_bundesland(db_pfad):
    r = runner.invoke(cli.app, ["export", "--bundesland", "XX", "--db", str(db_pfad)])
    assert r.exit_code == 2
    assert "nicht in settings.yaml" in r.output


@pytest.fixture
def fake_openplz(monkeypatch):
    from test_orte import openplz

    original = cli.OpenPlzOrte
    client, anfragen = openplz()
    monkeypatch.setattr(cli, "OpenPlzOrte", lambda url, **kw: original(url, **{**kw, "page_size": 5, "pause_s": 0},
                                                                         client=client))
    return anfragen


def _orte_args(db_pfad, tmp_path, *rest):
    gv = str(FIXTURES / "gv_isys" / "ausschnitt_sh.csv")
    return ["orte", "-b", "SH", "--db", str(db_pfad), "--gv", gv, *rest]


def test_orte_schreibt_neue_liste(fake_openplz, db_pfad, tmp_path):
    out = tmp_path / "orte_SH.csv"
    r = runner.invoke(cli.app, _orte_args(db_pfad, tmp_path, "--out", str(out)))
    assert r.exit_code == 0, r.output
    assert "8 Gemeinden (8 mit Einwohnerzahl), 5 Zusatz-Suchorte" in r.output
    assert "Neuenkirchen" in r.output  # Warnung: nur im GV-ISys
    assert "Suchumfang je Einwohnerstufe" in r.output and "Geschätzt" in r.output
    orte = load_orte(out)
    assert {o.name for o in orte if not o.ist_zusatz} >= {"Kiel", "Lübeck", "Heide", "Wesselburenerkoog"}
    # zweiter Lauf ohne --force bricht ab und lässt die Datei unverändert
    inhalt = out.read_text()
    out.write_text(inhalt + "# handbearbeitet\n")
    r = runner.invoke(cli.app, _orte_args(db_pfad, tmp_path, "--out", str(out)))
    assert r.exit_code == 2 and "existiert bereits" in r.output and "--force" in r.output
    assert out.read_text().endswith("# handbearbeitet\n")
    assert "Vergleich mit orte_SH.csv: bisher 9, neu 8 Gemeinden; 0 hinzu, 1 entfallen" in r.output  # Markierungszeile zählt mit
    r = runner.invoke(cli.app, _orte_args(db_pfad, tmp_path, "--out", str(out), "--force"))
    assert r.exit_code == 0 and out.read_text() == inhalt
    assert len(fake_openplz) == 3  # OpenPLZ-Seiten nur beim ersten Lauf geladen, danach aus dem Cache


def test_orte_bestehende_startliste_wird_nicht_ueberschrieben(fake_openplz, db_pfad, tmp_path, settings):
    startliste = settings.pfad(settings.bundesland("SH").orte_datei)
    vorher = startliste.read_bytes()
    r = runner.invoke(cli.app, _orte_args(db_pfad, tmp_path))
    assert r.exit_code == 2 and "existiert bereits" in r.output
    assert "bisher 80, neu 8 Gemeinden" in r.output and "entfallen (Namen prüfen)" in r.output
    assert startliste.read_bytes() == vorher


def test_orte_dry_run_und_ohne_gv(fake_openplz, db_pfad, tmp_path):
    out = tmp_path / "x.csv"
    r = runner.invoke(cli.app, ["orte", "-b", "SH", "--db", str(db_pfad), "--ohne-gv", "--dry-run", "--out", str(out)])
    assert r.exit_code == 0, r.output
    assert "8 Gemeinden (0 mit Einwohnerzahl), 0 Zusatz-Suchorte" in r.output and "nichts geschrieben" in r.output
    assert not out.exists()


def test_orte_gv_datei_fehlt(fake_openplz, db_pfad, tmp_path):
    r = runner.invoke(cli.app, ["orte", "-b", "SH", "--db", str(db_pfad), "--gv", str(tmp_path / "fehlt.xlsx")])
    assert r.exit_code == 2 and "GV-ISys-Datei" in r.output and "Destatis" in r.output


def test_orte_warnt_bei_kreisen_ohne_referenz(fake_openplz, db_pfad, tmp_path):
    r = runner.invoke(cli.app, _orte_args(db_pfad, tmp_path, "--dry-run"))
    assert "Kreise nicht in referenz_kammern.yaml" not in r.output  # alle Kreise der Testdaten sind dort


@pytest.fixture
def fake_web(monkeypatch):
    from fakes import FakeWeb, beispiel_seiten

    seiten = {**beispiel_seiten("foerde-steuer.example"), **beispiel_seiten("steuerbuero-schuett.example")}
    seiten["https://steuerbuero-schuett.example/"] = (301, "", {"location": "https://www.steuerbuero-schuett.example/index.php"})
    web = FakeWeb(seiten)
    original = cli.run_crawl
    schnell = lambda st, cfg, *a, **kw: original(st, cfg.model_copy(update={"requests_per_second_per_domain": 1000.0}),  # noqa: E731
                                                 *a, client=web.client(), **kw)
    monkeypatch.setattr(cli, "run_crawl", schnell)
    return web


@pytest.fixture
def fake_llm(monkeypatch):
    from test_extract_lauf import AufzeichnungsExtractor

    from recherche.normalize import PlzPruefung

    class Pruefer:
        anfragen = 0

        def __init__(self, *a, **kw):
            pass

        def pruefe(self, plz, ort):
            return PlzPruefung(True, "Schleswig-Holstein")

    monkeypatch.setattr(cli, "_extractor", lambda ctx, modell: AufzeichnungsExtractor())
    monkeypatch.setattr(cli, "PlzPruefer", Pruefer)


def test_crawl_extract_export(fake_web, fake_llm, db_pfad, tmp_path):
    with Storage(db_pfad) as st:
        for d in ("foerde-steuer.example", "steuerbuero-schuett.example", "weg.example"):
            st.add_kandidat(Kandidat(bundesland="SH", domain=d, website=f"https://{d}/", telefon="0461 1" if d == "weg.example" else None))
        st.add_kandidat(Kandidat(bundesland="SH", name="Kanzlei Ohneweb", plz="24103"))
    r = runner.invoke(cli.app, _args("crawl", db_pfad))
    assert r.exit_code == 0, r.output
    assert "Fertig: 2 Websites gelesen, 1 nicht gefunden" in r.output
    assert "1 × Startseite: HTTP 404" in r.output
    r = runner.invoke(cli.app, _args("extract", db_pfad))
    assert r.exit_code == 0, r.output
    assert "Fertig: 3 Kanzleien (1 ohne Website), 0 Fehler, 2 LLM-Aufrufe" in r.output
    assert "im Mittel 29.7 s und 1924 Token je Kanzlei" in r.output  # (2210+318 + 1105+214) / 2
    out = tmp_path / "x.xlsx"
    assert runner.invoke(cli.app, _args("export", db_pfad, "--out", str(out))).exit_code == 0
    wb = load_workbook(out)
    namen = {row[0].value: row for row in wb["Kanzleien"].iter_rows(min_row=2)}
    assert set(namen) == {"Förde Steuerberatungsgesellschaft mbH", "Steuerbüro Thomas Schütt", None}
    ng = list(wb["Nicht gefunden"].iter_rows(min_row=2, values_only=True))
    assert ng[0][0] == "weg.example" and "0461 1" in ng[0]
    protokoll = [r[1] for r in wb["Lauf-Protokoll"].iter_rows(min_row=2, values_only=True)]
    assert protokoll == ["crawl", "extract"]


def test_crawl_ohne_kandidaten(db_pfad):
    r = runner.invoke(cli.app, _args("crawl", db_pfad))
    assert r.exit_code == 0 and "keine offenen Kandidaten" in r.output


def test_extract_ohne_llm_und_wiederholen(fake_llm, db_pfad):
    r = runner.invoke(cli.app, _args("extract", db_pfad, "--ohne-llm", "--wiederholen"))
    assert r.exit_code == 0, r.output
    assert "0 fehlgeschlagene Kandidaten wieder offen" in r.output and "nur Regeln" in r.output


def test_extract_ollama_nicht_erreichbar(monkeypatch, fake_llm, db_pfad):
    from test_extract_lauf import AufzeichnungsExtractor, _kandidat_mit_seiten

    from recherche.extract.base import AnbieterNichtErreichbar

    with Storage(db_pfad) as st:
        _kandidat_mit_seiten(st, "steuerbuero-schuett.example")
    monkeypatch.setattr(cli, "_extractor", lambda ctx, m: AufzeichnungsExtractor(AnbieterNichtErreichbar("Ollama nicht erreichbar")))
    r = runner.invoke(cli.app, _args("extract", db_pfad))
    assert r.exit_code == 0 and "Lauf vorzeitig beendet: Ollama nicht erreichbar" in r.output


@pytest.fixture
def fake_openplz_pruefer(monkeypatch):
    from fakes import fake_pruefer

    monkeypatch.setattr(cli, "PlzPruefer", lambda *a, **kw: fake_pruefer())


def test_fixture_eval(fake_openplz_pruefer):
    r = runner.invoke(cli.app, ["fixture-eval", "--nur", "beispiele"])
    assert r.exit_code == 0, r.output
    assert "Fixtures: 6 ausgewertet; Modell qwen3.5:9b" in r.output
    assert "Trefferquote je Feld" in r.output and "6/6 (100%)" in r.output
    assert "Je Kanzlei im Mittel: 33.1 s, Token 3154 (Prompt 2842, Antwort 313) über 6 Antworten" in r.output
    assert "Modellvergleich" in r.output and "qwen3.6:27b" in r.output
    assert "✗" not in r.output.split("Modellvergleich")[0]


def test_fixture_eval_offline_ohne_ort_vereinheitlichung():
    r = runner.invoke(cli.app, ["fixture-eval", "--nur", "beispiele", "--offline"])
    assert r.exit_code == 0 and "nordlicht-steuerberatung.example · ort: ist 'Kiel-Holtenau' / soll 'Kiel'" in r.output


def test_fixture_eval_live_nur_echte(monkeypatch, fake_openplz_pruefer):
    aufrufe = []

    class Fake:
        def __init__(self, cfg, modell=None):
            self.modell = modell

        def extract(self, seiten):
            aufrufe.append(next(iter(seiten.values())).domain)
            raise AssertionError("darf für Beispiele nicht laufen")

    monkeypatch.setattr(cli, "OllamaExtractor", Fake)
    r = runner.invoke(cli.app, ["fixture-eval", "--live"])
    assert "--live: nur echte Fixtures" in r.output
    assert aufrufe == []  # echt/ ist hier leer, Beispiele werden nicht live überschrieben


def test_fixture_prompt():
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "nordlicht-steuerberatung.example"])
    assert r.exit_code == 0, r.output
    for teil in ("Zeichen je Bereinigungsstufe", "101 KB", "Linkerkennung", "/kanzlei/team/",
                 "Berufsträger 3", "2 Vertretungsberechtigte, Berufsbezeichnung: Steuerberater",
                 "geschätzt", "Zeichen/Token", "→ 3.05 Zeichen/Token"):
        assert teil in r.output, teil
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "mittelstand-steuer.example"])
    assert "Mitarbeiterzahl (Regel): 28 Expert:innen" in r.output and "/information/license" in r.output


def test_fixture_prompt_namen_und_versteckte_quellen():
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "nordlicht-steuerberatung.example"])
    assert "Teamseiten: Text außerhalb" in r.output and "Anna Kruse" not in r.output  # ohne --namen keine Namen
    assert "Berufsträger laut Zusammenführung (mit Aufzeichnung qwen3.5:9b): 3" in r.output
    assert "9/12 Felder belegt, 7 Personen" in r.output
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "nordlicht-steuerberatung.example", "--namen"])
    assert "(Anna Kruse, Lars Behrens)" in r.output and "Erster Ansprechpartner: Anna Kruse" in r.output
    assert "„Vertretungsberechtigte Partner:“ → 2 Namen in den Folgezeilen" in r.output
    assert "Alle Personen nach Zusammenführung (11)" in r.output and "| Teamseite, LLM | ja" in r.output and "Birte Claussen | Rechtsanwaltsfachangestellte | Teamseite | nein" in r.output


def test_fixture_prompt_live_und_speichern(monkeypatch, tmp_path):
    from test_extract_ollama import _ok, _ollama

    original = cli.OllamaExtractor
    client, gesendet = _ollama([_ok()])
    monkeypatch.setattr(cli, "OllamaExtractor", lambda cfg, modell=None: original(cfg, modell=modell, client=client))
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "hansen-steuerberatung.example", "--live", "--speichern"])
    assert r.exit_code == 0, r.output
    assert "Ollama qwen3.5:9b: 1200 Prompt-Token" in r.output and "Zeichen/Token" in r.output
    assert gesendet[0]["options"]["num_ctx"] == 16384
    ziel = load_settings().pfad("data") / "prompt_hansen-steuerberatung.example.txt"
    assert ziel.exists() and "=== SYSTEM ===" in ziel.read_text()
    ziel.unlink()


def test_fixture_prompt_unbekannt():
    r = runner.invoke(cli.app, ["fixture-prompt", "--domain", "gibtsnicht.example"])
    assert r.exit_code == 2 and "nicht gefunden" in r.output


def test_fixture_eval_ohne_aufzeichnung(fake_openplz_pruefer):
    r = runner.invoke(cli.app, ["fixture-eval", "--nur", "beispiele", "--modell", "gemma4:26b"])
    assert r.exit_code == 0
    assert "ohne LLM-Aufzeichnung" in r.output and "--live" in r.output
    assert "Je Kanzlei im Mittel" not in r.output


def test_fixture_eval_ohne_fixtures(tmp_path):
    r = runner.invoke(cli.app, ["fixture-eval", "--nur", "echt"])
    assert r.exit_code == 2 and "Keine Fixtures" in r.output


def test_fixture_add(monkeypatch, tmp_path):
    aufrufe = []

    def holen(domain, cfg, ziel):
        aufrufe.append((domain, ziel))
        (ziel / domain).mkdir(parents=True)
        (ziel / domain / "impressum.html").write_text("x")
        return ziel / domain, None

    monkeypatch.setattr(cli, "fixture_holen_sync", holen)
    r = runner.invoke(cli.app, ["fixture-add", "--domain", "Kanzlei-Beispiel.de", "--ziel", str(tmp_path)])
    assert r.exit_code == 0, r.output
    assert aufrufe == [("kanzlei-beispiel.de", tmp_path)]
    assert "impressum.html" in r.output and "erwartet.json" in r.output
    standard = load_settings().pfad(load_settings().pfade.fixtures_echt)
    runner.invoke(cli.app, ["fixture-add", "--domain", "b.example"])
    assert aufrufe[1][1] == standard
    import shutil

    shutil.rmtree(standard / "b.example")


def test_search_stufe_a_mit_kostenlog_und_wiederaufnahme(fakes, db_pfad):
    args = _args("search", db_pfad, "--orte", "Kiel,Husum")
    r = runner.invoke(cli.app, args)
    assert r.exit_code == 0, r.output
    assert "Websuche (Stufe A) mit brave" in r.output
    assert "Kostenlog brave: 8 Anfragen ≈ 0.04" in r.output  # 2 Orte x 4 Vorlagen
    with Storage(db_pfad) as st:
        assert [k.domain for k in st.kandidaten("SH")] == ["a.example", "h.example"]
    n = len(fakes["brave"].aufrufe)
    r = runner.invoke(cli.app, args)
    assert r.exit_code == 0
    assert len(fakes["brave"].aufrufe) == n
    assert "8 schon erledigt" in r.output


def test_search_abbruch_bei_limit(fakes, db_pfad):
    from recherche.search.base import KontingentErschoepft

    fakes["brave"].fehler_bei = {"Steuerberater Kiel": KontingentErschoepft("brave: Monatskontingent erschöpft")}
    r = runner.invoke(cli.app, _args("search", db_pfad, "--orte", "Kiel"))
    assert r.exit_code == 0
    assert "Lauf vorzeitig beendet" in r.output and "Monatskontingent" in r.output


def test_search_adapter_nicht_bereit(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("search", db_pfad, "--adapter", "google_cse"))
    assert r.exit_code == 2
    assert "GOOGLE_CSE_API_KEY fehlt" in r.output


def test_search_lehnt_lueckenfueller_ab(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("search", db_pfad, "--adapter", "places"))
    assert r.exit_code == 2 and "recherche fill --quelle places" in r.output


def test_fill_verlangt_ortsauswahl(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "places"))
    assert r.exit_code == 2 and "--orte" in r.output and "--alle" in r.output
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "places", "--orte", "Kiel", "--alle"))
    assert r.exit_code == 2


def test_fill_mit_orten(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "places", "--orte", "Kiel"))
    assert r.exit_code == 0, r.output
    assert "Lückenfüller (Stufe B) mit places in 1 Orten" in r.output
    assert fakes["places"].aufrufe == ["Kiel"]
    with Storage(db_pfad) as st:
        assert st.kandidaten("SH")[0].name == "Kanzlei Ohneweb"


def test_fill_lehnt_websuche_ab(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "brave", "--alle"))
    assert r.exit_code == 2 and "recherche search --adapter brave" in r.output


def test_fill_deaktivierte_quelle(db_pfad):
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "stb_verzeichnis", "--alle"))
    assert r.exit_code == 2 and "deaktiviert" in r.output


def _coverage_config(tmp_path):
    """settings.yaml-Kopie mit Test-Ortsliste und -Referenz."""
    cfg = tmp_path / "config"
    cfg.mkdir()
    s = load_settings()
    text = (s.basis_dir / "config" / "settings.yaml").read_text(encoding="utf-8")
    text = text.replace("orte_datei: config/orte_SH.csv", f"orte_datei: {FIXTURES / 'coverage' / 'orte.csv'}")
    text = text.replace("referenz: config/referenz_kammern.yaml", f"referenz: {FIXTURES / 'coverage' / 'referenz.yaml'}")
    text = text.replace("queries: config/queries.yaml", f"queries: {s.basis_dir / 'config' / 'queries.yaml'}")
    text = text.replace("blocklist: config/domain_blocklist.txt", f"blocklist: {s.basis_dir / 'config' / 'domain_blocklist.txt'}")
    (cfg / "settings.yaml").write_text(text, encoding="utf-8")
    return cfg / "settings.yaml"


def _coverage_daten(db_pfad):
    with Storage(db_pfad) as st:
        for i, k in enumerate(json.loads((FIXTURES / "coverage" / "kandidaten.json").read_text())["kandidaten"]):
            quellen = k.pop("quellen")
            kid = st.add_kandidat(Kandidat(bundesland="SH", **k))
            for q in quellen:
                st.add_fund(Fund(kandidat_id=kid, bundesland="SH", quelle=q, quell_id=f"{q}-{i}"))


def test_coverage(db_pfad, tmp_path):
    _coverage_daten(db_pfad)
    out = tmp_path / "cov.xlsx"
    r = runner.invoke(cli.app, _args("coverage", db_pfad, "--config", str(_coverage_config(tmp_path)), "--out", str(out)))
    assert r.exit_code == 0, r.output
    assert "19 Kanzleien gefunden, erwartet 50" in r.output
    assert "nur in einer Quelle: 16 (84%)" in r.output
    assert "Kreisfreie Stadt Kiel" in r.output and "Orte mit der größten Lücke" in r.output
    assert "Orte für `fill --luecken` (7): Kiel, Eckernförde" in r.output
    assert out.exists()


def test_coverage_ohne_referenz(db_pfad):
    r = runner.invoke(cli.app, _args("coverage", db_pfad))
    assert r.exit_code == 0
    assert "Keine (vollständigen) Referenzzahlen" in r.output


def test_fill_luecken(fakes, db_pfad, tmp_path):
    _coverage_daten(db_pfad)
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "places", "--luecken",
                                     "--config", str(_coverage_config(tmp_path))))
    assert r.exit_code == 0, r.output
    assert fakes["places"].aufrufe == ["Kiel", "Eckernförde", "Rendsburg", "Kronshagen", "Niebüll", "Husum", "Sylt"]


def test_fill_luecken_ohne_referenz(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("fill", db_pfad, "--quelle", "places", "--luecken"))
    assert r.exit_code == 2 and "Keine Orte unter der Referenz" in r.output


def test_run(fakes, fake_web, fake_llm, db_pfad, tmp_path):
    out = tmp_path / "run.xlsx"
    r = runner.invoke(cli.app, _args("run", db_pfad, "--orte", "Kiel", "--out", str(out)))
    assert r.exit_code == 0, r.output
    wb = load_workbook(out)
    assert wb["Lauf-Protokoll"].max_row > 1
    # a.example ist im Fake-Web nicht erreichbar → „Nicht gefunden" statt „Kanzleien"
    assert wb["Kanzleien"].max_row == 1
    assert wb["Nicht gefunden"]["A2"].value == "a.example"


def test_run_mit_ddgs(fakes, fake_web, fake_llm, db_pfad, tmp_path):
    r = runner.invoke(cli.app, _args("run", db_pfad, "--orte", "Kiel", "--adapter", "ddgs", "--out", str(tmp_path / "r.xlsx")))
    assert r.exit_code == 0, r.output
    assert "Websuche (Stufe A) mit ddgs" in r.output
    assert fakes["brave"].aufrufe == []


def test_search_compare_alle_aktiven(fakes, db_pfad, tmp_path):
    out = tmp_path / "vergleich.csv"
    r = runner.invoke(cli.app, _args("search-compare", db_pfad, "--orte", "Kiel,Husum", "--out", str(out)))
    assert r.exit_code == 0, r.output
    assert "übersprungen: google_cse – GOOGLE_CSE_API_KEY" in r.output
    assert "Überschneidung je Paar" in r.output
    assert "nur in ddgs: g.example" in r.output
    assert "nur in places" not in r.output  # Stufe B nur mit --adapter (kostet je Rasterzelle)
    assert out.exists()
    with Storage(db_pfad) as st:
        assert st.kandidaten("SH") == []  # Vergleich schreibt keine Kandidaten


def test_search_compare_auswahl(fakes, db_pfad):
    r = runner.invoke(cli.app, _args("search-compare", db_pfad, "--orte", "Kiel", "--adapter", "brave,ddgs"))
    assert r.exit_code == 0, r.output
    assert "places" not in r.output.split("Je Adapter")[1].split("Überschneidung")[0]


def test_extract_alle_mit_rueckfrage(tmp_path):
    from recherche.models import Kandidat, LaufEintrag, SchrittStatus
    from recherche.storage import Storage

    db = tmp_path / "t.sqlite"
    with Storage(db) as st:
        for i in range(3):
            st.add_kandidat(Kandidat(bundesland="SH", name=f"Kanzlei {i}", extract_status=SchrittStatus.FERTIG))
        st.log_lauf(LaufEintrag(lauf_id="l", schritt="extract", bundesland="SH", adapter="nur Regeln",
                                trefferzahl=10, laufzeit_s=20.0))
    r = runner.invoke(cli.app, ["extract", "--db", str(db), "-b", "SH", "--alle", "--ohne-llm"], input="n\n")
    assert r.exit_code == 0, r.output
    assert "3 Kandidaten" in r.output and "2.0 s je Kanzlei" in r.output and "Abgebrochen" in r.output
    with Storage(db) as st:
        assert all(k.extract_status == SchrittStatus.FERTIG for k in st.kandidaten("SH"))
    r = runner.invoke(cli.app, ["extract", "--db", str(db), "-b", "SH", "--alle", "--ohne-llm", "--ja"])
    assert r.exit_code == 0, r.output
    assert "3 Kandidaten wieder offen" in r.output and "Fertig: 3 Kanzleien" in r.output
