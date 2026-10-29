"""Standorte (SHBB-Niederlassungen) und Dubletten-Verdacht über Telefon + Ansprechpartner."""

from openpyxl import load_workbook

from recherche.export import BLATT_KANZLEIEN, export_excel
from recherche.extract.kanzlei import run_pruefung
from recherche.extract.standorte import dubletten_verdacht, standort_spalte
from recherche.models import Ansprechpartner, Kanzlei


def _k(domain, telefon, person, ort="Kiel", name=None, n_ap=1, **extra):
    aps = [Ansprechpartner(name=person, funktion="Steuerberater")] + [
        Ansprechpartner(name=f"Person {i}", funktion="Steuerberater") for i in range(n_ap - 1)]
    return Kanzlei(lauf_bundesland="SH", domain=domain, telefon=telefon, ort=ort, kanzleiname=name or domain,
                   ansprechpartner=aps, berufstraeger_anzahl=n_ap, **extra)


SHBB = [
    _k("buedelsdorf.shbb.example", "+49 431 59360", "Maik Jochens", ort="Büdelsdorf", name="SHBB Büdelsdorf"),
    _k("ihre-steuerberater.example", "+49 431 59360", "Maik Jochens", ort="Kiel", name="SHBB Zentrale", n_ap=4),
    _k("heide.shbb.example", "+49 481 1234", "Anna Meier", ort="Heide", name="SHBB Heide"),
    _k("shbb.example", "+49 431 59360", "Hans Lund", name="SHBB", n_ap=2),
    _k("andere-kanzlei.example", "+49 431 59360", "Eva Kühl"),  # gleiche Nummer, andere Person: nicht verbunden
]


def test_standorte_gruppiert():
    spalte = standort_spalte(SHBB)
    assert spalte["shbb.example"] == "Zentrale (4 Standorte)"  # Stammdomain selbst
    assert spalte["heide.shbb.example"] == "Standort von SHBB"
    assert spalte["ihre-steuerberater.example"] == "Standort von SHBB"  # über Telefon + Ansprechpartner
    assert "andere-kanzlei.example" not in spalte


def test_dubletten_verdacht_nur_bei_verschiedenen_domains():
    (d,) = dubletten_verdacht(SHBB, "SH")
    assert {d.behalten, d.dublette} == {"buedelsdorf.shbb.example", "ihre-steuerberater.example"}
    assert d.kriterium == "telefon+ansprechpartner" and d.art.value == "verdacht"
    assert "Büdelsdorf" in d.details and "Maik Jochens" in d.details


def test_export_und_pruefung(storage, tmp_path):
    for k in SHBB:
        storage.upsert_kanzlei(k)
    assert run_pruefung(storage, "SH", kanzlei_erkennung=False).dubletten == 1
    run_pruefung(storage, "SH", kanzlei_erkennung=False)
    assert len(storage.dubletten("SH")) == 1  # erneuter Lauf legt nichts doppelt an
    ws = load_workbook(export_excel(storage, "SH", tmp_path / "x.xlsx"))[BLATT_KANZLEIEN]
    kopf = [c.value for c in ws[1]]
    werte = {r[kopf.index("Domain")]: r[kopf.index("Zentrale/Standort")] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert werte["shbb.example"] == "Zentrale (4 Standorte)" and werte["andere-kanzlei.example"] is None
