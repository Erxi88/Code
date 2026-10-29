import json

import pytest
from conftest import FIXTURES
from openpyxl import load_workbook

from recherche.coverage import UNBEKANNT, als_excel, berechne, load_referenz
from recherche.models import Fund, Kandidat, Kanzlei, Status
from recherche.search.pipeline import load_orte

COV = FIXTURES / "coverage"
RD = "Kreis Rendsburg-Eckernförde"
NF = "Kreis Nordfriesland"
KIEL = "Kreisfreie Stadt Kiel"


@pytest.fixture
def orte():
    return load_orte(COV / "orte.csv")


@pytest.fixture
def referenz():
    return load_referenz(COV / "referenz.yaml", "SH")


@pytest.fixture
def daten(storage):
    for i, k in enumerate(json.loads((COV / "kandidaten.json").read_text(encoding="utf-8"))["kandidaten"]):
        quellen = k.pop("quellen")
        kid = storage.add_kandidat(Kandidat(bundesland="SH", quelle=quellen[0], **k))
        for q in quellen:
            storage.add_fund(Fund(kandidat_id=kid, bundesland="SH", quelle=q, quell_id=f"{q}-{i}"))
    return storage


def test_referenz_laden(referenz):
    assert (referenz.einheit, referenz.faktor, referenz.gesamt) == ("mitglieder", 0.5, 100)
    assert referenz.erwartet_gesamt == 50
    assert referenz.umrechnen(60) == 30
    assert referenz.kreise[RD] is None


def test_referenz_projektdatei_ist_gueltig(settings):
    # lokal mit echten Zahlen gefüllt – hier nur prüfen, dass die Datei lesbar ist
    r = load_referenz(settings.pfad(settings.pfade.referenz), "SH")
    assert r is not None and r.einheit in ("kanzleien", "mitglieder")


def test_referenz_ungueltige_einheit(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text("SH:\n  einheit: personen\n", encoding="utf-8")
    with pytest.raises(ValueError, match="einheit"):
        load_referenz(p, "SH")


def test_gefunden_je_kreis(daten, orte, referenz):
    erg = berechne(daten, "SH", orte, referenz)
    assert erg.gefunden == 19
    assert {k.kreis: k.gefunden for k in erg.kreise} == {KIEL: 10, RD: 4, NF: 4, UNBEKANNT: 1}


def test_erwartet_je_kreis(daten, orte, referenz):
    erg = berechne(daten, "SH", orte, referenz)
    kiel, rd, nf = erg.kreis(KIEL), erg.kreis(RD), erg.kreis(NF)
    assert (kiel.erwartet, kiel.basis, kiel.luecke) == (30, "kreis", 20)
    assert kiel.quote == pytest.approx(1 / 3)
    # Rest 50 - 30 = 20, verteilt nach Einwohnern 62.000 : 33.000 (Sylt ohne Einwohnerzahl)
    assert rd.erwartet == pytest.approx(20 * 62 / 95) and rd.basis == "verteilt"
    assert nf.erwartet == pytest.approx(20 * 33 / 95)
    assert erg.kreis(UNBEKANNT).erwartet is None
    assert erg.erwartet == 50


def test_quellen_und_nur_eine_quelle(daten, orte, referenz):
    erg = berechne(daten, "SH", orte, referenz)
    q = {z.quelle: (z.kanzleien, z.nur_diese) for z in erg.quellen}
    assert q == {"brave": (16, 13), "places": (4, 2), "gelbeseiten": (2, 1)}
    assert erg.nur_eine_quelle == 16
    assert erg.anteil_nur_eine_quelle == pytest.approx(16 / 19)


def test_orte_und_luecken(daten, orte, referenz):
    erg = berechne(daten, "SH", orte, referenz)
    o = {z.ort: z for z in erg.orte}
    assert (o["Kiel"].gefunden, o["Kiel"].erwartet) == (10, 30)
    assert o["Rendsburg"].gefunden == 3  # inkl. Websuche-Treffer ohne Adresse (Suchort)
    assert o["Sylt"].gefunden == 1 and o["Sylt"].erwartet is None  # Zuordnung über PLZ
    assert o["Kronshagen"].erwartet == pytest.approx(20 * 62 / 95 * 12 / 62)
    assert o["Eckernförde"].luecke == pytest.approx(20 * 62 / 95 * 22 / 62 - 1)
    # Orte mit eigener Erwartung nach Lücke, danach Orte ohne Einwohner in Lücken-Kreisen
    assert erg.luecken_orte(0.9) == ["Kiel", "Eckernförde", "Rendsburg", "Kronshagen", "Niebüll", "Husum", "Sylt"]
    assert "Husum" not in erg.luecken_orte(0.6)  # 3 >= 4,84 * 0,6


def test_ohne_referenz(daten, orte):
    erg = berechne(daten, "SH", orte, None)
    assert erg.erwartet is None
    assert all(k.erwartet is None for k in erg.kreise)
    assert erg.luecken_orte(0.9) == []
    assert erg.gefunden == 19


def test_nur_kreiswerte_ohne_gesamt(daten, orte, tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text("SH:\n  einheit: kanzleien\n  kreise:\n    Kreisfreie Stadt Kiel: 12\n", encoding="utf-8")
    erg = berechne(daten, "SH", orte, load_referenz(p, "SH"))
    assert erg.kreis(KIEL).erwartet == 12
    assert erg.kreis(RD).erwartet is None  # nichts zu verteilen
    assert erg.luecken_orte(0.9) == ["Kiel"]


def test_excel(daten, orte, referenz, tmp_path):
    pfad = als_excel(berechne(daten, "SH", orte, referenz), tmp_path / "cov.xlsx")
    wb = load_workbook(pfad)
    assert wb.sheetnames == ["Kreise", "Orte", "Quellen", "Zählung"]
    assert wb["Kreise"]["A1"].value == "Kreis"
    assert wb["Orte"].max_row == 8
    assert wb["Quellen"].max_row == 4


def test_neben_plz_und_zusatz_suchorte_zaehlen_zur_gemeinde(storage):
    from recherche.search.pipeline import Ort

    orte = [Ort("Kiel", "24103", "Kreisfreie Stadt Kiel", 247717, plz_weitere=("24105", "24159")),
            Ort("Kiel (24159)", "24159", "Kreisfreie Stadt Kiel", gemeinde="Kiel"),
            Ort("Heide", "25746", "Kreis Dithmarschen", 21540)]
    storage.add_kandidat(Kandidat(bundesland="SH", name="A", plz="24105"))  # Neben-PLZ
    storage.add_kandidat(Kandidat(bundesland="SH", name="B", suchort="Kiel (24159)"))  # Zusatz-Suchort
    storage.add_kandidat(Kandidat(bundesland="SH", name="C", ort="Heide", plz="25746"))
    erg = berechne(storage, "SH", orte, None)
    assert {o.ort: o.gefunden for o in erg.orte} == {"Kiel": 2, "Heide": 1}  # Zusatzort keine eigene Zeile
    assert erg.kreis(KIEL).gefunden == 2 and UNBEKANNT not in {k.kreis for k in erg.kreise}


def test_nach_extraktion_nur_kanzleien_im_land(storage):
    """Pilot: 1.169 Kandidaten, aber nur echte SH-Kanzleien zählen; Kreis aus der extrahierten PLZ."""
    from recherche.search.pipeline import Ort

    orte = [Ort("Kiel", "24103", "Kreisfreie Stadt Kiel", 247717), Ort("Heide", "25746", "Kreis Dithmarschen", 21540)]
    ids = [storage.add_kandidat(Kandidat(bundesland="SH", domain=f"k{i}.example", suchort="Kiel")) for i in range(6)]
    for i, (status, land, plz, ort) in enumerate([
        (Status.NEU, "Schleswig-Holstein", "25746", "Heide"),  # Suchort Kiel, Kanzlei aber in Heide
        (Status.GEPRUEFT, "Schleswig-Holstein", "24103", "Kiel"),
        (Status.KEINE_KANZLEI, "Schleswig-Holstein", "24103", "Kiel"),
        (Status.AUSSERHALB, "Sachsen", "01067", "Dresden"),
        (Status.NEU, None, None, None),  # PLZ nicht geprüft
    ]):  # fmt: skip
        storage.upsert_kanzlei(Kanzlei(kandidat_id=ids[i], domain=f"k{i}.example", lauf_bundesland="SH",
                                       status=status, bundesland=land, plz=plz, ort=ort))
    erg = berechne(storage, "SH", orte, None, ziel_land="Schleswig-Holstein")
    assert (erg.modus, erg.gefunden, erg.offen) == ("kanzleien", 2, 1)  # k5 ohne Kanzlei-Zeile = offen
    assert (erg.keine_kanzlei, erg.ausserhalb, erg.ohne_bundesland) == (1, 1, 1)
    assert {o.ort: o.gefunden for o in erg.orte} == {"Kiel": 1, "Heide": 1}
    # vor der Extraktion: Kandidaten zählen (fill --luecken nach search)
    leer_storage_erg = berechne(storage, "HH", orte, None, ziel_land="Hamburg")
    assert leer_storage_erg.modus == "kandidaten"
