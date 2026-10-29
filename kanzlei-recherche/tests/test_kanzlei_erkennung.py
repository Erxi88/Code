"""Kanzlei-Erkennung: Magazine, Branchenbücher und Verlage mit sauberem Impressum sind keine Kanzleien."""

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from recherche import cli
from recherche.blocklist_vorschlag import vorschlaege
from recherche.export import BLATT_KANZLEIEN, BLATT_KEINE_KANZLEI, export_excel
from recherche.extract.base import LlmErgebnis
from recherche.extract.kanzlei import (
    anwenden,
    grund_aus_abweichung,
    kein_kanzlei_grund,
    run_pruefung,
)
from recherche.extract.lauf import regeln_fuer
from recherche.extract.merge import zusammenfuehren
from recherche.extract.rules import Ausschluss
from recherche.models import (
    Fund,
    Kandidat,
    Kanzlei,
    Konfidenz,
    Seite,
    Seitentyp,
    Status,
)
from recherche.normalize import PlzPruefung


def _seite(domain, typ, zeilen):
    html = "<html><body>" + "".join(f"<p>{z}</p>" for z in zeilen) + "</body></html>"
    return Seite(url=f"https://{domain}/{typ.value}", domain=domain, seitentyp=typ, http_status=200,
                 text="\n".join(zeilen), html=html)


def _verifiziert(plz, ort):
    return PlzPruefung(verifiziert=True, bundesland="Schleswig-Holstein", ort_kanonisch=ort)


def _fund(domain, ort="Kiel", name=None, snippet=None, i=0):
    return Fund(kandidat_id=1, bundesland="SH", quelle="brave", quell_id=f"{domain}-{i}", suchort=ort,
                name=name, snippet=snippet, website=f"https://{domain}/")


def _kanzlei(domain, impressum, team=None, llm_name=None, funde=()):
    seiten = {Seitentyp.IMPRESSUM: _seite(domain, Seitentyp.IMPRESSUM, impressum)}
    if team:
        seiten[Seitentyp.TEAM] = _seite(domain, Seitentyp.TEAM, team)
    k = zusammenfuehren(Kandidat(id=1, bundesland="SH", domain=domain), list(funde), seiten,
                        regeln_fuer(seiten, Ausschluss(), domain),
                        LlmErgebnis(kanzleiname=llm_name), _verifiziert)
    return k, seiten


MAGAZIN_IMPRESSUM = [
    "Impressum", "Kiel Magazin Verlag GmbH", "Holstenstraße 1", "24103 Kiel", "Telefon: 0431 1234560",
    "E-Mail: redaktion@kiel-magazin.example", "Geschäftsführer: Hans Meier",
    "Registergericht: Amtsgericht Kiel, HRB 12345", "Verantwortlich für den Inhalt nach § 18 Abs. 2 MStV: Hans Meier",
]


def test_magazin_mit_sauberem_verlagsimpressum_ist_keine_kanzlei():
    from recherche.extract.kanzlei import pruefen

    funde = [_fund("kiel-magazin.example", name="Die besten Steuerberater in Kiel", snippet="Unsere Auswahl …")]
    k, seiten = _kanzlei("kiel-magazin.example", MAGAZIN_IMPRESSUM, llm_name="Kiel Magazin Verlag GmbH", funde=funde)
    assert k.konfidenz == Konfidenz.HOCH  # ohne Prüfung: sauberes Impressum, PLZ verifiziert
    # Titel nennt Steuerberater – zählt hier nur, weil der Treffer aus einem Ort kommt; aus vielen Orten nicht
    viele = [_fund("kiel-magazin.example", ort=f"Ort {i}", name="Die besten Steuerberater", i=i) for i in range(5)]
    pruefen(k, seiten, funde=viele)
    assert k.status == Status.KEINE_KANZLEI and k.konfidenz == Konfidenz.MITTEL
    grund = grund_aus_abweichung(k.abweichung)
    assert grund.startswith("keine Steuerberaterkammer") and "5 Suchorten (Portal)" in grund
    # ohne Suchtreffer genauso
    pruefen(k, seiten)
    assert k.status == Status.KEINE_KANZLEI and grund_aus_abweichung(k.abweichung).endswith("kein Steuerberater im Suchtreffer")
    assert k.abweichung.count("Keine Kanzlei:") == 1  # Grund wird ersetzt, nicht angehängt


def test_kanzlei_nur_ueber_kammer_im_impressum():
    impressum = ("Impressum\nGodbersen GmbH\nHauptstraße 5\n25813 Husum\nZuständige Aufsichtsbehörde:\n"
                 "Steuerberaterkammer Schleswig-Holstein, Hopfenstraße 2d, 24114 Kiel")
    k = Kanzlei(lauf_bundesland="SH", domain="godbersen.example", kanzleiname="Godbersen GmbH", rechtsform="GmbH")
    assert kein_kanzlei_grund(k, impressum) is None
    assert kein_kanzlei_grund(k, impressum.replace("Steuerberaterkammer", "Industrie- und Handelskammer")) is not None


def test_kanzlei_nur_mit_berufstraeger_auf_teamseite():
    impressum = ["Impressum", "Godbersen GmbH", "Hauptstraße 5", "25813 Husum", "Telefon: 04841 123450",
                 "Geschäftsführer: Jan Godbersen", "Registergericht: Amtsgericht Flensburg, HRB 999"]
    team = ["Unser Team", "Jan Godbersen", "Steuerberater", "Maike Lund", "Steuerfachangestellte"]
    k, seiten = _kanzlei("godbersen.example", impressum, team=team, llm_name="Godbersen GmbH")
    assert k.rechtsform == "GmbH" and k.berufstraeger_anzahl == 1
    anwenden(k, kein_kanzlei_grund(k, seiten[Seitentyp.IMPRESSUM].text), impressum_gelesen=True)
    assert k.status == Status.NEU and k.konfidenz == Konfidenz.HOCH
    # ohne Teamseite fehlt das Merkmal
    k2, seiten2 = _kanzlei("godbersen.example", impressum, llm_name="Godbersen GmbH")
    assert k2.berufstraeger_anzahl == 0
    assert kein_kanzlei_grund(k2, seiten2[Seitentyp.IMPRESSUM].text) is not None


@pytest.mark.parametrize("name, snippet", [
    ("Spieckermann | Goetz", "Ihre Steuerberater in Kiel – persönlich und digital"),
    ("Spieckermann | Goetz – Steuerkanzlei", None),
    ("Spieckermann | Goetz", "StB Spieckermann, Kiel"),
])  # fmt: skip
def test_kanzlei_nur_ueber_snippet(name, snippet):
    k = Kanzlei(lauf_bundesland="SH", domain="spieckermann-goetz.example", kanzleiname="Spieckermann | Goetz")
    funde = [_fund("spieckermann-goetz.example", name=name, snippet=snippet)]
    assert kein_kanzlei_grund(k, "Impressum\nSpieckermann | Goetz\nKiel") is not None
    assert kein_kanzlei_grund(k, "Impressum\nSpieckermann | Goetz\nKiel", funde=funde) is None


@pytest.mark.parametrize("name, rechtsform", [
    ("Treuhand Nord GmbH", "GmbH"), ("Muster Revision GmbH", "GmbH"), ("Muster WP GmbH Wirtschaftsprüfungsgesellschaft", "GmbH"),
    ("Kanzlei am Wall", None), ("Meier & Lund", "PartG mbB"), ("Hans Meier", "Einzelkanzlei"),
])  # fmt: skip
def test_name_oder_rechtsform_genuegt(name, rechtsform):
    assert kein_kanzlei_grund(Kanzlei(lauf_bundesland="SH", kanzleiname=name, rechtsform=rechtsform), None) is None


def _speichern(st, domain, name, impressum, funde=(), status=Status.NEU, konfidenz=Konfidenz.HOCH):
    kid = st.add_kandidat(Kandidat(bundesland="SH", domain=domain, website=f"https://{domain}/"))
    for i, f in enumerate(funde):
        st.add_fund(f.model_copy(update={"kandidat_id": kid, "quell_id": f"{domain}-{i}"}))
    seite = _seite(domain, Seitentyp.IMPRESSUM, impressum)
    st.save_seite(seite)
    st.setze_kandidat_seite(kid, Seitentyp.IMPRESSUM, seite.url)
    st.upsert_kanzlei(Kanzlei(kandidat_id=kid, domain=domain, lauf_bundesland="SH", kanzleiname=name, plz="24103",
                              ort="Kiel", plz_verifiziert=True, konfidenz=konfidenz, status=status))
    return kid


def test_pruefung_bestand_excel_und_blocklist(storage, tmp_path):
    orte = [_fund("x", ort=f"Ort {i}", name="Branchenbuch Kiel", i=i) for i in range(6)]
    _speichern(storage, "kiel-magazin.example", None, MAGAZIN_IMPRESSUM, funde=orte)
    _speichern(storage, "kanzlei-am-wall.example", "Kanzlei am Wall", ["Impressum", "Kanzlei am Wall"])
    _speichern(storage, "branchenbuch.example", None, ["Impressum", "Branchenbuch GmbH"], status=Status.GEPRUEFT)
    stat = run_pruefung(storage, "SH")
    assert (stat.geprueft, stat.keine_kanzlei, stat.herabgestuft) == (3, 2, 2)
    nach = {k.domain: k for k in storage.kanzleien("SH")}
    assert nach["kiel-magazin.example"].status == Status.KEINE_KANZLEI
    assert nach["kiel-magazin.example"].konfidenz == Konfidenz.NIEDRIG  # Kernfelder fehlen
    assert nach["kanzlei-am-wall.example"].status == Status.NEU
    assert nach["branchenbuch.example"].status == Status.GEPRUEFT  # von Hand gepflegt, bleibt

    pfad = export_excel(storage, "SH", tmp_path / "x.xlsx")
    wb = load_workbook(pfad)
    kanzleien = [r[0] for r in wb[BLATT_KANZLEIEN].iter_rows(min_row=2, values_only=True)]
    keine = list(wb[BLATT_KEINE_KANZLEI].iter_rows(values_only=True))
    assert len(kanzleien) == 2 and len(keine) == 2  # Kopf + Magazin, mit allen Spalten
    kopf = keine[0]
    assert kopf[0] == "Grund" and "Kanzleiname" in kopf and "Telefon" in kopf and "Domain" in kopf
    assert keine[1][kopf.index("Domain")] == "kiel-magazin.example" and "Steuerberaterkammer" in keine[1][0]
    assert "keine Kanzlei" not in wb[BLATT_KANZLEIEN].data_validations.dataValidation[0].formula1

    liste = vorschlaege(storage, "SH")
    assert [(v.domain, v.grund_kurz, v.portal) for v in liste] == [("kiel-magazin.example", "keine Kanzlei", True)]

    # Kanzlei kommt zurück, wenn ein Merkmal hinzukommt (z. B. Kammer im neu gelesenen Impressum)
    kid = nach["kiel-magazin.example"].kandidat_id
    seite = _seite("kiel-magazin.example", Seitentyp.IMPRESSUM, [*MAGAZIN_IMPRESSUM, "Steuerberaterkammer Schleswig-Holstein"])
    seite.url += "-neu"
    storage.save_seite(seite)
    storage.setze_kandidat_seite(kid, Seitentyp.IMPRESSUM, seite.url)
    assert run_pruefung(storage, "SH").wieder_kanzlei == 1
    k = next(k for k in storage.kanzleien("SH") if k.domain == "kiel-magazin.example")
    assert k.status == Status.NEU and grund_aus_abweichung(k.abweichung) is None


def test_ohne_excelblatt_wenn_keine(storage, tmp_path):
    _speichern(storage, "kanzlei-am-wall.example", "Kanzlei am Wall", ["Impressum"])
    assert BLATT_KEINE_KANZLEI not in load_workbook(export_excel(storage, "SH", tmp_path / "x.xlsx")).sheetnames


def test_cli(tmp_path, monkeypatch):
    from recherche.normalize import PlzPruefung
    from recherche.storage import Storage

    class OhneNetz:  # kein OpenPLZ in Tests
        def __init__(self, *a, **k):
            pass

        def pruefe(self, plz, ort):
            return PlzPruefung(None)

    monkeypatch.setattr(cli, "PlzPruefer", OhneNetz)

    db = tmp_path / "t.sqlite"
    with Storage(db) as st:
        _speichern(st, "kiel-magazin.example", None, MAGAZIN_IMPRESSUM)
    r = CliRunner().invoke(cli.app, ["kanzlei-pruefung", "--db", str(db), "-b", "SH"])
    assert r.exit_code == 0, r.output
    assert "1 keine Kanzlei" in r.output and "kiel-magazin.example" in r.output


def test_portal_mit_vielen_subdomains_zaehlt_je_stammdomain(storage):
    """city-map.de: viele Kandidaten (Subdomains je Ort) mit je einem Suchort und Steuerberater-Titel."""
    for i in range(6):
        domain = f"ort{i}.city-map.example"
        _speichern(storage, domain, None, ["Impressum", "City Map GmbH", "Geschäftsführer: Max Muster"],
                   funde=[_fund(domain, ort=f"Ort {i}", name=f"Steuerberater in Ort {i} – City Map")])
    _speichern(storage, "kanzlei-am-wall.example", "Kanzlei am Wall", ["Impressum"],
               funde=[_fund("kanzlei-am-wall.example", ort="Kiel")])
    stat = run_pruefung(storage, "SH")
    assert stat.keine_kanzlei == 6
    k = next(k for k in storage.kanzleien("SH") if k.domain == "ort0.city-map.example")
    assert "Domain in 6 Suchorten (Portal)" in grund_aus_abweichung(k.abweichung)


def test_ab_fuenf_suchorten_zaehlen_nur_kammer_und_berufstraeger():
    k = Kanzlei(lauf_bundesland="SH", kanzleiname="Steuerberatung Nord GmbH", rechtsform="PartG mbB")
    assert kein_kanzlei_grund(k, None, suchorte=4) is None
    assert "Portal" in kein_kanzlei_grund(k, None, suchorte=5)
    assert kein_kanzlei_grund(k, "Steuerberaterkammer Schleswig-Holstein", suchorte=78) is None
    assert kein_kanzlei_grund(k.model_copy(update={"berufstraeger_anzahl": 2}), None, suchorte=78) is None


@pytest.mark.parametrize("name, rechtsform", [
    (None, "Einzelkanzlei"),  # Portal-Impressum mit einem Vertreter → Einzelkanzlei
    ("11880-steuerberater.com", None),
    ("Steuerberater in Kiel – jetzt vergleichen", None),
    ("Die besten Steuerberater in Husum", "GmbH"),
])  # fmt: skip
def test_ohne_echten_kanzleinamen_keine_kanzlei(name, rechtsform):
    k = Kanzlei(lauf_bundesland="SH", kanzleiname=name, rechtsform=rechtsform)
    funde = [_fund("x.example", name="Steuerberater in Kiel")]
    assert "kein Kanzleiname" in kein_kanzlei_grund(k, None, funde=funde)


def test_stammdomain_baukasten_hoster_bleiben_getrennt():
    from recherche.extract.kanzlei import stammdomain

    assert stammdomain("kiel.city-map.de") == "city-map.de"
    assert stammdomain("www.kanzlei.example") == "kanzlei.example"
    assert stammdomain("holgerherholz.wixsite.com") == "holgerherholz.wixsite.com"
    assert stammdomain("steuer-nord.jimdofree.com") == "steuer-nord.jimdofree.com"


def test_ausserhalb_bundesland_eigenes_blatt_und_vorrang(storage, tmp_path):
    from recherche.export import BLATT_AUSSERHALB

    for domain, land, name in [("kanzlei-dresden.example", "Sachsen", "Steuerkanzlei Dresden"),
                               ("kanzlei-kiel.example", "Schleswig-Holstein", "Steuerkanzlei Kiel"),
                               ("magazin-leipzig.example", "Sachsen", None)]:
        kid = _speichern(storage, domain, name, ["Impressum", name or "Verlag GmbH"])
        k = storage.kanzlei_fuer_kandidat(kid)
        storage.upsert_kanzlei(k.model_copy(update={"bundesland": land}))
    # ohne Website (Verzeichnis): nur Bundesland-Filter
    storage.upsert_kanzlei(Kanzlei(lauf_bundesland="SH", kanzleiname="Steuerbüro Hamburg", bundesland="Hamburg",
                                   website_vorhanden=False))
    stat = run_pruefung(storage, "SH", ziel_land="Schleswig-Holstein")
    status = {k.kanzleiname or k.domain: k.status for k in storage.kanzleien("SH")}
    assert status == {"Steuerkanzlei Dresden": Status.AUSSERHALB, "Steuerkanzlei Kiel": Status.NEU,
                      "magazin-leipzig.example": Status.KEINE_KANZLEI,  # keine Kanzlei geht vor
                      "Steuerbüro Hamburg": Status.AUSSERHALB}
    assert stat.ausserhalb == 2
    wb = load_workbook(export_excel(storage, "SH", tmp_path / "x.xlsx"))
    aussen = list(wb[BLATT_AUSSERHALB].iter_rows(min_row=2, values_only=True))
    kanzleien = [r[0] for r in wb[BLATT_KANZLEIEN].iter_rows(min_row=2, values_only=True)]
    assert len(aussen) == 2 and kanzleien == ["Steuerkanzlei Kiel"], kanzleien
    # Zurück, wenn die PLZ korrigiert wird; von Hand gepflegter Status bleibt
    k = next(k for k in storage.kanzleien("SH") if k.kanzleiname == "Steuerkanzlei Dresden")
    storage.upsert_kanzlei(k.model_copy(update={"bundesland": "Schleswig-Holstein"}))
    k = next(k for k in storage.kanzleien("SH") if k.kanzleiname == "Steuerbüro Hamburg")
    storage.upsert_kanzlei(k.model_copy(update={"status": Status.GEPRUEFT}))
    run_pruefung(storage, "SH", ziel_land="Schleswig-Holstein")
    status = {k.kanzleiname or k.domain: k.status for k in storage.kanzleien("SH")}
    assert status["Steuerkanzlei Dresden"] == Status.NEU and status["Steuerbüro Hamburg"] == Status.GEPRUEFT
