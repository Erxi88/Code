"""Coverage: gefundene Kanzleien je Kreis, Ort und Quelle gegen Referenzzahlen.

Gezählt werden – sobald extrahiert ist – nur Kanzleien mit Status „Kanzlei" (nicht „keine Kanzlei",
nicht „außerhalb Bundesland") und Bundesland laut PLZ = Bundesland des Laufs; Kreis und Ort aus der
extrahierten Adresse. Kandidaten ohne Kanzlei-Zeile (nicht gelesen) stehen getrennt als „offen".
Vor der ersten Extraktion (nach `search`, für `fill --luecken`) zählen die Kandidaten selbst.

Zuordnung zu Ort/Kreis (beste verfügbare Angabe): Ort (Adresse) → PLZ → Suchort → „unbekannt".
Vor dem Crawl haben reine Websuche-Treffer keine Adresse; dann zählt der Suchort,
was unscharf ist (die Suche „Steuerberater Kiel" findet auch Umlandkanzleien).

Erwartete Zahlen (config/referenz_kammern.yaml):
- je Kreis aus eigenem Wert, sonst Rest von `gesamt`, verteilt nach Einwohnern der Orte
  (Spalte einwohner der Ortsliste),
- je Ort: Kreiswert anteilig nach Einwohnern der gelisteten Orte des Kreises,
- bei einheit: mitglieder wird mit `faktor` in Kanzleien umgerechnet.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

from .models import SchrittStatus, Status
from .search.pipeline import Ort
from .storage import Storage

UNBEKANNT = "unbekannt"


@dataclass
class Referenz:
    bundesland: str
    einheit: str = "kanzleien"  # kanzleien | mitglieder
    faktor: float | None = None
    gesamt: float | None = None
    kreise: dict[str, float | None] = field(default_factory=dict)
    quelle: str | None = None
    stand: str | None = None

    def umrechnen(self, wert: float | None) -> float | None:
        """Referenzwert → erwartete Kanzleien (None, wenn nicht umrechenbar)."""
        if wert is None:
            return None
        if self.einheit == "mitglieder":
            return wert * self.faktor if self.faktor is not None else None
        return float(wert)

    @property
    def erwartet_gesamt(self) -> float | None:
        if self.gesamt is not None:
            return self.umrechnen(self.gesamt)
        werte = [self.umrechnen(w) for w in self.kreise.values()]
        return sum(werte) if werte and all(w is not None for w in werte) else None


def load_referenz(pfad: str | Path, bundesland: str) -> Referenz | None:
    pfad = Path(pfad)
    if not pfad.exists():
        return None
    daten = (yaml.safe_load(pfad.read_text(encoding="utf-8-sig")) or {}).get(bundesland.upper())
    if not daten:
        return None
    einheit = daten.get("einheit") or "kanzleien"
    if einheit not in ("kanzleien", "mitglieder"):
        raise ValueError(f"referenz {bundesland}: einheit muss 'kanzleien' oder 'mitglieder' sein, nicht {einheit!r}")
    return Referenz(
        bundesland=bundesland.upper(),
        einheit=einheit,
        faktor=daten.get("faktor"),
        gesamt=daten.get("gesamt"),
        kreise=dict(daten.get("kreise") or {}),
        quelle=daten.get("quelle"),
        stand=str(daten["stand"]) if daten.get("stand") else None,
    )


@dataclass
class KreisZeile:
    kreis: str
    gefunden: int
    erwartet: float | None = None
    basis: str | None = None  # "kreis" (eigener Wert) | "verteilt" (Rest von gesamt)

    @property
    def quote(self) -> float | None:
        return self.gefunden / self.erwartet if self.erwartet else None

    @property
    def luecke(self) -> float | None:
        return self.erwartet - self.gefunden if self.erwartet is not None else None


@dataclass
class OrtZeile:
    ort: str
    kreis: str | None
    gefunden: int
    einwohner: int | None = None
    erwartet: float | None = None

    @property
    def luecke(self) -> float | None:
        return self.erwartet - self.gefunden if self.erwartet is not None else None


@dataclass
class QuellenZeile:
    quelle: str
    kanzleien: int  # Kandidaten mit mindestens einem Fund dieser Quelle
    nur_diese: int  # … die keine andere Quelle hat

    @property
    def anteil_nur_diese(self) -> float:
        return self.nur_diese / self.kanzleien if self.kanzleien else 0.0


@dataclass
class CoverageErgebnis:
    bundesland: str
    gefunden: int
    erwartet: float | None
    kreise: list[KreisZeile]
    orte: list[OrtZeile]
    quellen: list[QuellenZeile]
    nur_eine_quelle: int
    referenz: Referenz | None = None
    modus: str = "kandidaten"  # kanzleien (nach der Extraktion) | kandidaten (vorher)
    offen: int = 0  # Kandidaten ohne Kanzlei-Zeile (noch nicht oder nicht lesbar), ohne ausserhalb_vorab
    weitere_domains: int = 0  # weitere Domains derselben Kanzlei – einmal gezählt (wie in der Versandliste)
    zu_pruefen: int = 0  # im Bundesland, aber im Blatt „Zu prüfen" (Prüfkatalog Q5/Q6/Q12, 1B Rang 4/5)
    ausserhalb_vorab: int = 0  # nicht gecrawlt: Suchtreffer nennen nur PLZ außerhalb (Vorfilter vor dem Crawl)
    keine_kanzlei: int = 0
    ausserhalb: int = 0
    ohne_bundesland: int = 0  # PLZ nicht geprüft/unbekannt – nicht gezählt

    @property
    def anteil_nur_eine_quelle(self) -> float:
        return self.nur_eine_quelle / self.gefunden if self.gefunden else 0.0

    def kreis(self, name: str) -> KreisZeile:
        return next(k for k in self.kreise if k.kreis == name)

    def luecken_orte(self, quote_ziel: float) -> list[str]:
        """Orte unter der Referenz, größte Lücke zuerst (Eingabe für `fill --luecken`).

        Orte mit eigener Erwartung: gefunden < erwartet * quote_ziel. Orte ohne
        Einwohnerzahl: alle Orte eines Kreises, der selbst unter der Referenz liegt.
        """
        kreise = {k.kreis: k for k in self.kreise}
        mit: list[tuple[float, str]] = []
        ohne: list[tuple[float, str]] = []
        for o in self.orte:
            if o.erwartet is not None:
                if o.gefunden < o.erwartet * quote_ziel:
                    mit.append((o.luecke or 0.0, o.ort))
                continue
            k = kreise.get(o.kreis or "")
            if k and k.erwartet is not None and k.gefunden < k.erwartet * quote_ziel:
                ohne.append((k.luecke or 0.0, o.ort))
        return [o for _, o in sorted(mit, reverse=True)] + [o for _, o in sorted(ohne, reverse=True)]


def _zuordnen(storage: Storage, bundesland: str, orte: list[Ort],
              eintraege: list[tuple[int, str | None, str | None, str | None]] | None = None) -> dict[int, Ort | None]:
    """Kandidat → Gemeinde. Zusatz-Suchorte (PLZ-Bezirke großer Städte) und Neben-PLZ
    werden ihrer Gemeinde zugeordnet. `eintraege`: (Kandidat-ID, Ort, PLZ, Suchort) statt der Kandidaten."""
    gemeinden = [o for o in orte if not o.ist_zusatz]
    nach_name = {o.name.casefold(): o for o in gemeinden}
    for z in orte:
        if z.ist_zusatz and z.gemeinde and z.gemeinde.casefold() in nach_name:
            nach_name.setdefault(z.name.casefold(), nach_name[z.gemeinde.casefold()])
    nach_plz: dict[str, Ort] = {}
    for o in gemeinden:
        for plz in o.alle_plz:
            nach_plz.setdefault(plz, o)
    if eintraege is None:
        eintraege = [(k.id, k.ort, k.plz, k.suchort) for k in storage.kandidaten(bundesland)]
    zuordnung: dict[int, Ort | None] = {}
    for kid, k_ort, k_plz, suchort in eintraege:
        ort = None
        if k_ort:
            ort = nach_name.get(k_ort.casefold())
        if ort is None and k_plz:
            ort = nach_plz.get(k_plz)
        if ort is None and k_ort is None and k_plz is None and suchort:
            # automatisch nachgesuchte Stadtteile („Kiel (24105)") gehören zur Gemeinde
            ort = nach_name.get(suchort.casefold()) or nach_name.get(re.sub(r"\s*\(\d{5}\)$", "", suchort).casefold())
        zuordnung[kid] = ort
    return zuordnung


def berechne(
    storage: Storage, bundesland: str, orte: list[Ort], referenz: Referenz | None = None,
    ziel_land: str | None = None,
) -> CoverageErgebnis:
    """`ziel_land` („Schleswig-Holstein") gesetzt und schon extrahiert → nur echte Kanzleien im Land zählen."""
    kanzleien = storage.kanzleien(bundesland) if ziel_land else []
    zusatz: dict[str, int] = {}
    if kanzleien:
        from .extract.standorte import spiegel_domains

        # weitere Domains derselben Kanzlei (Knüppel, wetreu …) zählen wie in der Versandliste nur einmal
        weitere = {id(k) for liste in spiegel_domains(kanzleien).values() for k in liste}
        from .qualitaet import zu_pruefen_grund

        # wie das Hauptblatt: ohne weitere Domains derselben Kanzlei und ohne „Zu prüfen" (Prüfkatalog Q14)
        gezaehlt = [k for k in kanzleien if k.kandidat_id is not None and k.status not in
                    (Status.KEINE_KANZLEI, Status.AUSSERHALB) and k.bundesland == ziel_land and id(k) not in weitere
                    and zu_pruefen_grund(k) is None]
        mit_zeile = {k.kandidat_id for k in kanzleien}
        ohne_zeile = [k for k in storage.kandidaten(bundesland) if k.id not in mit_zeile]
        vorab = sum(1 for k in ohne_zeile if k.crawl_status == SchrittStatus.FEHLER and "laut Suchtreffer" in (k.fehler or ""))
        zusatz = dict(
            modus="kanzleien",
            offen=len(ohne_zeile) - vorab,
            ausserhalb_vorab=vorab,
            keine_kanzlei=sum(1 for k in kanzleien if k.status == Status.KEINE_KANZLEI),
            ausserhalb=sum(1 for k in kanzleien if k.status == Status.AUSSERHALB),
            ohne_bundesland=sum(1 for k in kanzleien if k.status not in (Status.KEINE_KANZLEI, Status.AUSSERHALB)
                                and k.bundesland != ziel_land),
            weitere_domains=sum(1 for k in kanzleien if id(k) in weitere and k.bundesland == ziel_land),
            zu_pruefen=sum(1 for k in kanzleien if id(k) not in weitere and k.bundesland == ziel_land
                           and zu_pruefen_grund(k) is not None),
        )
        suchorte = {k.id: k.suchort for k in storage.kandidaten(bundesland)}
        zuordnung = _zuordnen(storage, bundesland, orte,
                              [(k.kandidat_id, k.ort, k.plz, suchorte.get(k.kandidat_id)) for k in gezaehlt])
    else:
        zuordnung = _zuordnen(storage, bundesland, orte)
    orte = [o for o in orte if not o.ist_zusatz]  # Erwartung und Lücken nur je Gemeinde

    # ---- gefunden je Kreis und Ort
    je_kreis: dict[str, int] = defaultdict(int)
    je_ort: dict[str, int] = defaultdict(int)
    for ort in zuordnung.values():
        je_kreis[(ort.kreis if ort else None) or UNBEKANNT] += 1
        if ort:
            je_ort[ort.name] += 1

    # ---- Quellen
    quellen_je_kandidat: dict[int, set[str]] = defaultdict(set)
    for f in storage.funde(bundesland):
        if f.kandidat_id in zuordnung:
            quellen_je_kandidat[f.kandidat_id].add(f.quelle)
    alle_quellen = sorted({q for qs in quellen_je_kandidat.values() for q in qs})
    quellen = [
        QuellenZeile(
            q,
            sum(1 for qs in quellen_je_kandidat.values() if q in qs),
            sum(1 for qs in quellen_je_kandidat.values() if qs == {q}),
        )
        for q in alle_quellen
    ]
    nur_eine = sum(1 for kid in zuordnung if len(quellen_je_kandidat.get(kid, ())) == 1)

    # ---- erwartete Zahlen je Kreis
    kreis_namen = list(dict.fromkeys([o.kreis for o in orte if o.kreis] + list(referenz.kreise if referenz else [])))
    einwohner_kreis: dict[str, int] = defaultdict(int)
    for o in orte:
        if o.kreis and o.einwohner:
            einwohner_kreis[o.kreis] += o.einwohner
    erwartet: dict[str, tuple[float | None, str | None]] = {}
    if referenz:
        eigene = {k: referenz.umrechnen(w) for k, w in referenz.kreise.items() if referenz.umrechnen(w) is not None}
        rest = referenz.erwartet_gesamt
        if referenz.gesamt is not None and rest is not None:
            rest -= sum(eigene.values())
        ohne = [k for k in kreis_namen if k not in eigene]
        summe_ew = sum(einwohner_kreis.get(k, 0) for k in ohne)
        for k in kreis_namen:
            if k in eigene:
                erwartet[k] = (eigene[k], "kreis")
            elif referenz.gesamt is not None and rest is not None and summe_ew and einwohner_kreis.get(k):
                erwartet[k] = (max(rest, 0.0) * einwohner_kreis[k] / summe_ew, "verteilt")
    kreise = [KreisZeile(k, je_kreis.get(k, 0), *erwartet.get(k, (None, None))) for k in kreis_namen]
    if je_kreis.get(UNBEKANNT):
        kreise.append(KreisZeile(UNBEKANNT, je_kreis[UNBEKANNT]))

    # ---- Orte
    orte_zeilen = []
    for o in orte:
        kreis_erw = erwartet.get(o.kreis or "", (None, None))[0]
        ort_erw = None
        if kreis_erw is not None and o.einwohner and einwohner_kreis.get(o.kreis or ""):
            ort_erw = kreis_erw * o.einwohner / einwohner_kreis[o.kreis]
        orte_zeilen.append(OrtZeile(o.name, o.kreis, je_ort.get(o.name, 0), o.einwohner, ort_erw))

    return CoverageErgebnis(
        bundesland=bundesland,
        gefunden=len(zuordnung),
        erwartet=referenz.erwartet_gesamt if referenz else None,
        kreise=kreise,
        orte=orte_zeilen,
        quellen=quellen,
        nur_eine_quelle=nur_eine,
        referenz=referenz,
        **zusatz,
    )


def als_excel(erg: CoverageErgebnis, pfad: str | Path) -> Path:
    pfad = Path(pfad)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    blaetter = [
        ("Kreise", ["Kreis", "Gefunden", "Erwartet", "Basis", "Quote", "Lücke"],
         [[k.kreis, k.gefunden, _r(k.erwartet), k.basis, _r(k.quote, 3), _r(k.luecke)] for k in erg.kreise]),
        ("Orte", ["Ort", "Kreis", "Einwohner", "Gefunden", "Erwartet", "Lücke"],
         [[o.ort, o.kreis, o.einwohner, o.gefunden, _r(o.erwartet), _r(o.luecke)] for o in erg.orte]),
        ("Quellen", ["Quelle", "Kanzleien", "Nur in dieser Quelle", "Anteil nur diese"],
         [[q.quelle, q.kanzleien, q.nur_diese, _r(q.anteil_nur_diese, 3)] for q in erg.quellen]),
        ("Zählung", ["Kennzahl", "Anzahl"],
         [["gezählt als", erg.modus], ["gefunden (gezählt)", erg.gefunden], ["erwartet", _r(erg.erwartet)],
          ["keine Kanzlei", erg.keine_kanzlei], ["außerhalb Bundesland", erg.ausserhalb],
          ["ohne geprüfte PLZ", erg.ohne_bundesland], ["außerhalb laut Suchtreffer (nicht gecrawlt)", erg.ausserhalb_vorab],
          ["weitere Domains derselben Kanzlei (einmal gezählt)", erg.weitere_domains],
          ["zu prüfen (Blatt „Zu prüfen“)", erg.zu_pruefen],
          ["offen (nicht gelesen)", erg.offen]]),
    ]  # fmt: skip
    wb.remove(wb.active)
    for titel, kopf, zeilen in blaetter:
        ws = wb.create_sheet(titel)
        ws.append(kopf)
        for c in ws[1]:
            c.font = Font(bold=True)
        for z in zeilen:
            ws.append(z)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(pfad)
    return pfad


def _r(wert: float | None, stellen: int = 1) -> float | None:
    return round(wert, stellen) if wert is not None else None
