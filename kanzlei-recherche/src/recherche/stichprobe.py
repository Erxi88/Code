"""Stichprobe für die Abnahme vor jedem Volllauf (Prüfkatalog Abschnitt 5).

`recherche stichprobe -b SH -n 50 --seed 1` zieht n zufällige Zeilen des Hauptblatts (wie im Export: ohne „Keine
Kanzlei", „Außerhalb", „Zu prüfen" und weitere Domains) – mit festem Seed reproduzierbar – in eine Excel-Datei mit
Blatt „Stichprobe": je Kernfeld (Name, Straße, PLZ/Ort, Telefon, Ansprechpartner) der Wert und eine Prüfspalte mit
Auswahl richtig / falsch / leer, dazu „Kanzlei?" und „Ausland?" (ja/nein) und die Quell-URLs zum Nachsehen.

`recherche stichprobe --auswerten datei.xlsx` zählt die Prüfspalten: bestanden bei höchstens 2 falschen Werten je
Kernfeld in 50 Zeilen (≤ 4 %), null Zeilen, die keine Kanzlei sind, und null Ausland.
"""

from __future__ import annotations

import random
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Kanzlei
from .qualitaet import im_hauptblatt

BLATT = "Stichprobe"
KERNFELDER: list[tuple[str, callable]] = [
    ("Name", lambda k: k.kanzleiname),
    ("Straße", lambda k: " ".join(x for x in (k.strasse, k.adresszusatz) if x) or None),
    ("PLZ/Ort", lambda k: " ".join(x for x in (k.plz, k.ort) if x) or None),
    ("Telefon", lambda k: k.telefon),
    ("Ansprechpartner", lambda k: k.hauptansprechpartner.anzeige() if k.hauptansprechpartner else None),
]
PRUEF_WERTE = ("richtig", "falsch", "leer")
MAX_FALSCH_ANTEIL = 0.04  # ≤ 2 von 50


def ziehen(kanzleien: Sequence[Kanzlei], n: int, seed: int) -> list[Kanzlei]:
    """n zufällige Zeilen des Hauptblatts, reproduzierbar: feste Grundreihenfolge (Domain/Id), dann Seed."""
    from .extract.standorte import spiegel_domains

    weitere = {id(k) for liste in spiegel_domains(list(kanzleien)).values() for k in liste}
    haupt = sorted((k for k in kanzleien if im_hauptblatt(k, weitere)), key=lambda k: (k.domain or "", k.id or 0))
    return random.Random(seed).sample(haupt, min(n, len(haupt)))


def schreiben(liste: Sequence[Kanzlei], pfad: Path, bundesland: str, seed: int) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = BLATT
    kopf = ["Nr.", "Domain"]
    for feld, _ in KERNFELDER:
        kopf += [feld, f"{feld} – Prüfung"]
    kopf += ["Kanzlei?", "Ausland?", "Bemerkung", "Konfidenz", "Quell-URLs", "Website"]
    breiten = {"Nr.": 5, "Domain": 26, "Bemerkung": 30, "Konfidenz": 10, "Quell-URLs": 55, "Website": 30}
    for i, titel in enumerate(kopf, start=1):
        z = ws.cell(row=1, column=i, value=titel)
        z.font, z.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[get_column_letter(i)].width = breiten.get(
            titel, 13 if titel.endswith(("Prüfung", "?")) else 30)
    for r, k in enumerate(liste, start=2):
        werte = [r - 1, k.domain]
        for _, wert in KERNFELDER:
            werte += [wert(k), None]
        werte += [None, None, None, k.konfidenz.value, "\n".join(k.quell_urls) or None, k.website]
        for c, w in enumerate(werte, start=1):
            zelle = ws.cell(row=r, column=c, value=w)
            if kopf[c - 1] in ("Quell-URLs", "Bemerkung"):
                zelle.alignment = Alignment(wrap_text=True, vertical="top")
            if kopf[c - 1] in ("PLZ/Ort", "Telefon"):
                zelle.number_format = "@"
    letzte = len(liste) + 1
    for titel, werte in [(f"{f} – Prüfung", PRUEF_WERTE) for f, _ in KERNFELDER] + [("Kanzlei?", ("ja", "nein")),
                                                                                       ("Ausland?", ("nein", "ja"))]:
        b = get_column_letter(kopf.index(titel) + 1)
        dv = DataValidation(type="list", formula1=f'"{",".join(werte)}"', allow_blank=True)
        dv.add(f"{b}2:{b}{max(letzte, 2)}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kopf))}{max(letzte, 1)}"
    info = wb.create_sheet("Info")
    for zeile in (["Bundesland", bundesland], ["Seed", seed], ["Zeilen", len(liste)],
                  ["Abnahme", "höchstens 2 falsche Werte je Kernfeld in 50 Zeilen (≤ 4 %), 0 keine Kanzlei, 0 Ausland"],
                  ["Auswerten", f"recherche stichprobe --auswerten {pfad.name}"]):
        info.append(zeile)
    pfad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pfad)
    return pfad


@dataclass
class Auswertung:
    zeilen: int = 0
    falsch: dict[str, int] = field(default_factory=dict)
    leer: dict[str, int] = field(default_factory=dict)
    ungeprueft: dict[str, int] = field(default_factory=dict)
    keine_kanzlei: int = 0
    ausland: int = 0

    @property
    def grenze(self) -> int:
        """Erlaubte falsche Werte je Kernfeld: 4 % der Zeilen, abgerundet (50 → 2)."""
        return int(self.zeilen * MAX_FALSCH_ANTEIL + 1e-9)

    @property
    def gruende(self) -> list[str]:
        g = [f"{f}: {n} falsch (erlaubt {self.grenze})" for f, n in self.falsch.items() if n > self.grenze]
        if self.keine_kanzlei:
            g.append(f"{self.keine_kanzlei} Zeile(n) keine Kanzlei (erlaubt 0)")
        if self.ausland:
            g.append(f"{self.ausland} Zeile(n) Ausland (erlaubt 0)")
        if offen := sum(self.ungeprueft.values()):
            g.append(f"{offen} Prüfzelle(n) noch leer")
        return g

    @property
    def bestanden(self) -> bool:
        return self.zeilen > 0 and not self.gruende


def auswerten(pfad: Path) -> Auswertung:
    ws = load_workbook(pfad, read_only=True)[BLATT]
    zeilen = list(ws.iter_rows(values_only=True))
    kopf = list(zeilen[0])
    erg = Auswertung(falsch={f: 0 for f, _ in KERNFELDER}, leer={f: 0 for f, _ in KERNFELDER},
                     ungeprueft={f: 0 for f, _ in KERNFELDER})
    for z in zeilen[1:]:
        if not any(z):
            continue
        erg.zeilen += 1
        for f, _ in KERNFELDER:
            wert = str(z[kopf.index(f"{f} – Prüfung")] or "").strip().lower()
            if wert == "falsch":
                erg.falsch[f] += 1
            elif wert == "leer":
                erg.leer[f] += 1
            elif wert != "richtig":
                erg.ungeprueft[f] += 1
        erg.keine_kanzlei += str(z[kopf.index("Kanzlei?")] or "").strip().lower() == "nein"
        erg.ausland += str(z[kopf.index("Ausland?")] or "").strip().lower() == "ja"
    return erg
