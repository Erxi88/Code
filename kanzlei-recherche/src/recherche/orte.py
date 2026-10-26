"""Vollständige Ortsliste eines Bundeslandes (recherche orte).

1. OpenPLZ /de/FederalStates/<Land>/Localities – alle Seiten, jede Seite gecacht.
   Eine Zeile je (Ortsname, PLZ); zusammengefasst je Gemeinde über den amtlichen
   Gemeindeschlüssel (AGS, municipality.key): Name, alle PLZ, Kreis (Rückfall).
2. GV-ISys (Destatis, „Alle politisch selbständigen Gemeinden"), XLSX oder CSV:
   Einwohner und Kreisname je AGS. Kreis-Zeilen (Satzart 40) liefern den Kreisnamen,
   Gemeinde-Zeilen (Satzart 60) die Einwohner. Spaltennamen sind einstellbar.
3. Große kreisfreie Städte (Kiel, Lübeck) erhalten Zusatz-Suchorte je PLZ-Bezirk, weil
   „Steuerberater Kiel" nach drei Ergebnisseiten abbricht.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import time
from collections import defaultdict
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlsplit

import httpx

from .config import GvSpalten, StadtteilConfig
from .search.pipeline import Ort

log = logging.getLogger(__name__)

SATZART_KREIS = "40"
SATZART_GEMEINDE = "60"
# Textkennzeichen der Kreis-Zeilen → Bezeichnung (wie in der Startliste und referenz_kammern.yaml)
KREIS_TYPEN = {"41": "Kreisfreie Stadt", "42": "Stadtkreis", "43": "Kreis", "44": "Landkreis", "45": "Regionalverband"}


class OrteFehler(RuntimeError):
    """Ortsliste kann nicht erzeugt werden (Quelle fehlt, Spalte fehlt, OpenPLZ nicht erreichbar)."""


def gemeinde_name(name: str) -> str:
    """'Kiel, Landeshauptstadt' → 'Kiel'; 'Glücksburg (Ostsee), Stadt' → 'Glücksburg (Ostsee)'."""
    return name.split(",", 1)[0].strip()


def kreis_bezeichnung(name: str, typ: str | None) -> str:
    """'Dithmarschen' + 'Kreis' → 'Kreis Dithmarschen'; 'Kiel, Landeshauptstadt' + 'Kreisfreie Stadt'
    → 'Kreisfreie Stadt Kiel'."""
    name = gemeinde_name(name)
    if not typ or name.lower().startswith(typ.lower()):
        return name
    return f"{typ} {name}"


# ---------------------------------------------------------------------- OpenPLZ


class OpenPlzOrte:
    """Lädt alle Localities eines Bundeslandes seitenweise (gecacht in api_cache)."""

    name = "openplz_orte"

    def __init__(
        self,
        url: str = "https://openplzapi.org",
        cache=None,
        client: httpx.Client | None = None,
        page_size: int = 50,
        pause_s: float = 0.5,
        retries: int = 3,
        sleep: Callable[[float], None] = time.sleep,
    ):
        self.url = url.rstrip("/")
        self.cache = cache
        self._client = client
        self.page_size = page_size
        self.pause_s = pause_s
        self.retries = retries
        self._sleep = sleep
        self.anfragen = 0

    @property
    def client(self) -> httpx.Client:
        if self._client is None:
            self._client = httpx.Client(timeout=30.0)
        return self._client

    def localities(self, land_schluessel: str, fortschritt: Callable[[str], None] | None = None) -> list[dict]:
        """Alle Einträge. Ende: x-total-pages erreicht oder Seite kürzer als page_size."""
        alle: list[dict] = []
        seite = 1
        while seite <= 2000:  # Schutz gegen Endlosschleifen
            daten, gesamt = self._seite(land_schluessel, seite)
            alle.extend(daten)
            if fortschritt:
                fortschritt(f"OpenPLZ Seite {seite}{f'/{gesamt}' if gesamt else ''}: {len(daten)} Einträge")
            if not daten or len(daten) < self.page_size or (gesamt is not None and seite >= gesamt):
                break
            seite += 1
        return alle

    def _seite(self, land: str, seite: int) -> tuple[list[dict], int | None]:
        schluessel = f"FederalStates/{land}/Localities?page={seite}&pageSize={self.page_size}"
        if self.cache is not None and (gecacht := self.cache.cache_get(self.name, schluessel)) is not None:
            return gecacht["daten"], gecacht.get("gesamt_seiten")
        url = f"{self.url}/de/FederalStates/{land}/Localities"
        letzter = ""
        for versuch in range(1, self.retries + 2):
            if self.anfragen:
                self._sleep(self.pause_s)
            self.anfragen += 1
            try:
                resp = self.client.get(url, params={"page": seite, "pageSize": self.page_size})
            except httpx.HTTPError as e:
                letzter = f"{type(e).__name__}: {e}"
            else:
                if resp.status_code == 200:
                    gesamt = resp.headers.get("x-total-pages")
                    ergebnis = {"daten": resp.json(), "gesamt_seiten": int(gesamt) if gesamt and gesamt.isdigit() else None}
                    if self.cache is not None:
                        self.cache.cache_put(self.name, schluessel, ergebnis)
                    return ergebnis["daten"], ergebnis["gesamt_seiten"]
                letzter = f"HTTP {resp.status_code}"
                if resp.status_code < 500 and resp.status_code != 429:
                    break
            if versuch <= self.retries:
                self._sleep(2.0 * 2 ** (versuch - 1))
        raise OrteFehler(f"OpenPLZ Seite {seite} nicht abrufbar ({letzter})")


@dataclass
class Gemeinde:
    ags: str
    name: str
    plz: list[str] = field(default_factory=list)
    ortsnamen: dict[str, set[str]] = field(default_factory=lambda: defaultdict(set))  # PLZ → Ortsnamen
    kreis_openplz: str | None = None


def gemeinden_aus_localities(localities: Iterable[dict]) -> dict[str, Gemeinde]:
    gemeinden: dict[str, Gemeinde] = {}
    for loc in localities:
        m = loc.get("municipality") or {}
        ags = (m.get("key") or "").strip()
        plz = (loc.get("postalCode") or "").strip()
        if not ags:
            continue
        g = gemeinden.get(ags)
        if g is None:
            d = loc.get("district") or {}
            kreis = kreis_bezeichnung(d["name"], d.get("type")) if d.get("name") else None
            g = gemeinden[ags] = Gemeinde(ags, gemeinde_name(m.get("name") or loc.get("name") or ags), kreis_openplz=kreis)
        if plz and plz not in g.plz:
            g.plz.append(plz)
        if plz and loc.get("name"):
            g.ortsnamen[plz].add(loc["name"].strip())
    for g in gemeinden.values():
        g.plz.sort()
    return gemeinden


# ---------------------------------------------------------------------- GV-ISys


@dataclass
class GvDaten:
    einwohner: dict[str, int] = field(default_factory=dict)  # AGS → Einwohner
    namen: dict[str, str] = field(default_factory=dict)  # AGS → Gemeindename
    kreise: dict[str, str] = field(default_factory=dict)  # Kreisschlüssel (5 Stellen) → Bezeichnung
    kreis_je_ags: dict[str, str] = field(default_factory=dict)  # nur CSV mit Kreisname-Spalte
    kreisfrei: set[str] = field(default_factory=set)  # Kreisschlüssel kreisfreier Städte


def _norm(text) -> str:
    t = str(text or "").lower().replace("-\n", "").replace("­", "").replace("-", "")
    return " ".join(re.sub(r"[^a-z0-9äöüß]+", " ", t).split())


def _spalte_finden(labels: list[str], name: str | None) -> int | None:
    """Spaltenindex: exakter Name, dann Name als eigenes Wort, dann eindeutige Teilübereinstimmung."""
    if not name:
        return None
    n = _norm(name)
    for pruefung in (lambda lab: lab == n, lambda lab: n in lab.split() or lab.endswith(" " + n)):
        treffer = [i for i, lab in enumerate(labels) if lab and pruefung(lab)]
        if treffer:
            return treffer[0]
    teil = [i for i, lab in enumerate(labels) if lab and n in lab]
    return teil[0] if len(teil) == 1 else None


def _ziffern(wert, stellen: int) -> str | None:
    if wert is None or str(wert).strip() == "":
        return None
    t = re.sub(r"\D", "", str(wert).split(".")[0] if isinstance(wert, float) else str(wert))
    return t.zfill(stellen) if t else None


def _zahl(wert) -> int | None:
    if wert is None:
        return None
    if isinstance(wert, (int, float)):
        return int(wert)
    t = re.sub(r"[^\d]", "", str(wert))
    return int(t) if t else None


def _zeilen_xlsx(pfad: Path, blatt: str | None, merkmal: str) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(pfad, read_only=True, data_only=True)
    blaetter = [wb[blatt]] if blatt else wb.worksheets
    for ws in blaetter:
        zeilen = [list(r) for r in ws.iter_rows(values_only=True)]
        if any(_norm(merkmal) in _norm(c) for z in zeilen[:30] for c in z):
            return zeilen
    raise OrteFehler(f"Kein Blatt mit Spalte {merkmal!r} in {pfad.name}")


def _zeilen_csv(pfad: Path) -> list[list]:
    roh = pfad.read_bytes()
    for kodierung in ("utf-8-sig", "cp1252"):
        try:
            text = roh.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    trenner = ";" if text.count(";") >= text.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=trenner)]


def lade_gv(pfad: Path, spalten: GvSpalten, land_schluessel: str, blatt: str | None = None) -> GvDaten:
    """GV-ISys-Auszug (XLSX oder CSV) einlesen, nur Zeilen des Bundeslandes."""
    merkmal = spalten.satzart or spalten.name
    zeilen = _zeilen_xlsx(pfad, blatt, merkmal) if pfad.suffix.lower() in (".xlsx", ".xlsm") else _zeilen_csv(pfad)
    # Kopfzeile: erste Zeile mit dem Merkmal; mehrzeilige Köpfe bis zur ersten Datenzeile zusammenfassen
    kopf = next((i for i, z in enumerate(zeilen[:40]) if any(_norm(merkmal) in _norm(c) for c in z)), None)
    if kopf is None:
        raise OrteFehler(f"Spalte {merkmal!r} nicht gefunden – Spaltennamen in settings.yaml (orte.gv_isys.spalten) prüfen")
    breite = max(len(z) for z in zeilen[kopf : kopf + 4])
    labels = [""] * breite
    start = kopf + 1
    for i in range(kopf, min(kopf + 5, len(zeilen))):
        z = zeilen[i] + [None] * (breite - len(zeilen[i]))
        # erste Datenzeile: Schlüsselspalte (Satzart, sonst AGS, sonst Land) enthält eine Zahl
        schluessel = next((j for k in ("satzart", "ags", "land") if (j := _spalte_finden(labels, getattr(spalten, k))) is not None), None)
        if i > kopf and schluessel is not None and _ziffern(z[schluessel], 1):
            break
        labels = [f"{a} {_norm(b)}".strip() for a, b in zip(labels, z, strict=True)]
        start = i + 1
    idx = {k: _spalte_finden(labels, getattr(spalten, k)) for k in GvSpalten.model_fields}
    fehlend = [k for k in ("name", "einwohner") if idx[k] is None]
    if idx["ags"] is None and None in (idx["land"], idx["kreis"], idx["gem"]):
        fehlend.append("ags bzw. land/kreis/gem")
    if fehlend:
        raise OrteFehler(
            f"GV-ISys: Spalte(n) {', '.join(fehlend)} nicht gefunden. Gefundene Spalten: "
            f"{', '.join(lab for lab in labels if lab)} – Namen in settings.yaml (orte.gv_isys.spalten) anpassen"
        )

    def wert(z: list, k: str):
        i = idx[k]
        return z[i] if i is not None and i < len(z) else None

    gv = GvDaten()
    for z in zeilen[start:]:
        if idx["ags"] is not None and wert(z, "ags"):
            ags = _ziffern(wert(z, "ags"), 8)
            land, kreis_schl = (ags or "")[:2], (ags or "")[:5]
        else:
            land = _ziffern(wert(z, "land"), 2)
            rb = _ziffern(wert(z, "rb"), 1) or "0"
            kreis = _ziffern(wert(z, "kreis"), 2)
            gem = _ziffern(wert(z, "gem"), 3)
            kreis_schl = f"{land}{rb}{kreis}" if land and kreis else None
            ags = f"{kreis_schl}{gem}" if kreis_schl and gem else None
        if land != land_schluessel:
            continue
        satzart = _ziffern(wert(z, "satzart"), 2) if idx["satzart"] is not None else SATZART_GEMEINDE
        name = str(wert(z, "name") or "").strip()
        if satzart == SATZART_KREIS and kreis_schl and name:
            typ = KREIS_TYPEN.get(_ziffern(wert(z, "textkennzeichen"), 2) or "")
            gv.kreise[kreis_schl] = kreis_bezeichnung(name, typ)
            if typ in ("Kreisfreie Stadt", "Stadtkreis"):
                gv.kreisfrei.add(kreis_schl)
        elif satzart == SATZART_GEMEINDE and ags:
            if (ew := _zahl(wert(z, "einwohner"))) is not None:
                gv.einwohner[ags] = ew
            if name:
                gv.namen[ags] = gemeinde_name(name)
            if (kn := wert(z, "kreisname")) and str(kn).strip():
                gv.kreis_je_ags[ags] = str(kn).strip()
    if not gv.einwohner:
        raise OrteFehler(f"GV-ISys: keine Gemeinden für Land {land_schluessel} gefunden – Datei/Spalten prüfen")
    return gv


def gv_quelle_laden(quelle: str, ziel_dir: Path, client: httpx.Client | None = None) -> Path:
    """Pfad zurückgeben; eine URL wird einmal nach ziel_dir geladen (danach aus der Datei)."""
    if not re.match(r"^https?://", quelle):
        pfad = Path(quelle)
        if not pfad.exists():
            raise OrteFehler(
                f"GV-ISys-Datei {pfad} fehlt – bei Destatis (GV-ISys, „Alle politisch selbständigen Gemeinden“) "
                "herunterladen und Pfad/URL in settings.yaml (orte.gv_isys.quelle) eintragen oder --gv angeben"
            )
        return pfad
    name = Path(urlsplit(quelle).path).name or "gemeindeverzeichnis.xlsx"
    pfad = ziel_dir / name
    if pfad.exists():
        return pfad
    client = client or httpx.Client(timeout=120.0, follow_redirects=True)
    try:
        resp = client.get(quelle)
        resp.raise_for_status()
    except httpx.HTTPError as e:
        raise OrteFehler(f"GV-ISys-Download fehlgeschlagen ({e})") from e
    pfad.parent.mkdir(parents=True, exist_ok=True)
    pfad.write_bytes(resp.content)
    return pfad


# ---------------------------------------------------------------------- Zusammenführen


@dataclass
class OrteErgebnis:
    orte: list[Ort]
    warnungen: list[str] = field(default_factory=list)

    @property
    def gemeinden(self) -> list[Ort]:
        return [o for o in self.orte if not o.ist_zusatz]

    @property
    def zusatz(self) -> list[Ort]:
        return [o for o in self.orte if o.ist_zusatz]


def _kreis_kurz(kreis: str | None) -> str:
    return re.sub(r"^(Kreisfreie Stadt|Stadtkreis|Landkreis|Kreis|Regionalverband)\s+", "", kreis or "")


def orte_erzeugen(
    gemeinden: dict[str, Gemeinde], gv: GvDaten | None, stadtteile: StadtteilConfig | None = None
) -> OrteErgebnis:
    erg = OrteErgebnis([])
    zeilen: list[tuple[Ort, Gemeinde, bool]] = []
    for ags, g in gemeinden.items():
        kreis_schl = ags[:5]
        if gv is not None:
            kreis = gv.kreise.get(kreis_schl) or gv.kreis_je_ags.get(ags) or g.kreis_openplz
            einwohner = gv.einwohner.get(ags)
            name = gv.namen.get(ags) or g.name
            kreisfrei = kreis_schl in gv.kreisfrei or (kreis or "").startswith("Kreisfreie Stadt")
            if ags not in gv.einwohner:
                erg.warnungen.append(f"{g.name} ({ags}): nicht im GV-ISys – ohne Einwohnerzahl")
        else:
            kreis, einwohner, name = g.kreis_openplz, None, g.name
            kreisfrei = (kreis or "").startswith("Kreisfreie Stadt")
        ort = Ort(name=name, plz=g.plz[0] if g.plz else "", kreis=kreis, einwohner=einwohner, ags=ags,
                  plz_weitere=tuple(g.plz[1:]))  # fmt: skip
        zeilen.append((ort, g, kreisfrei))
    if gv is not None:
        for ags in sorted(set(gv.einwohner) - set(gemeinden)):
            erg.warnungen.append(f"{gv.namen.get(ags, ags)} ({ags}): im GV-ISys, aber nicht in OpenPLZ – nicht übernommen")

    # gleiche Namen in verschiedenen Kreisen eindeutig machen
    zaehler: dict[str, int] = defaultdict(int)
    for ort, _, _ in zeilen:
        zaehler[ort.name.casefold()] += 1
    eindeutig = []
    for ort, g, kreisfrei in zeilen:
        if zaehler[ort.name.casefold()] > 1:
            neu = f"{ort.name} ({_kreis_kurz(ort.kreis)})"
            erg.warnungen.append(f"Name {ort.name!r} mehrfach – als {neu!r} geführt")
            ort = Ort(**{**ort.__dict__, "name": neu})
        eindeutig.append((ort, g, kreisfrei))

    eindeutig.sort(key=lambda x: ((x[0].kreis or "~"), x[0].name.casefold()))
    for ort, g, kreisfrei in eindeutig:
        erg.orte.append(ort)
        st = stadtteile
        if st and st.aktiv and ort.einwohner and ort.einwohner >= st.ab_einwohner and (kreisfrei or not st.nur_kreisfrei):
            for plz in ort.alle_plz:
                eigene = sorted(n for n in g.ortsnamen.get(plz, ()) if _norm(gemeinde_name(n)) != _norm(ort.name))
                erg.orte.append(Ort(name=f"{ort.name} ({plz})", plz=plz, kreis=ort.kreis, ags=ort.ags,
                                    gemeinde=ort.name, stadtteil=eigene[0] if eigene else None))  # fmt: skip
    return erg


@dataclass
class Vergleich:
    alt: int
    neu: int
    hinzu: list[str]
    entfallen: list[str]
    mit_einwohnern: int
    zusatz: int


def vergleichen(alt: list[Ort], neu: list[Ort]) -> Vergleich:
    a = {o.name.casefold(): o.name for o in alt if not o.ist_zusatz}
    n = {o.name.casefold(): o.name for o in neu if not o.ist_zusatz}
    return Vergleich(
        alt=len(a), neu=len(n),
        hinzu=sorted(n[k] for k in n.keys() - a.keys()),
        entfallen=sorted(a[k] for k in a.keys() - n.keys()),
        mit_einwohnern=sum(1 for o in neu if not o.ist_zusatz and o.einwohner is not None),
        zusatz=sum(1 for o in neu if o.ist_zusatz),
    )  # fmt: skip


def _tsd(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def stufen_verteilung(orte: list[Ort], plan) -> dict[str, tuple[int, int, int]]:
    """Je Einwohnerstufe (in Reihenfolge der Regel): (Orte, Queries, Seiten je Query) – zur Kostenabschätzung."""
    namen: dict[int, str] = {}
    unten = 0
    for s in plan.regel:
        if s.unter is None:
            namen[id(s)] = f"ab {_tsd(unten)} Einwohner" if unten else "alle Orte"
        else:
            namen[id(s)] = f"unter {_tsd(s.unter)} Einwohner" if not unten else f"{_tsd(unten)}–{_tsd(s.unter - 1)} Einwohner"
            unten = s.unter
    reihenfolge = [namen[id(s)] for s in plan.regel] + ["ohne Einwohnerzahl", "Zusatz-Suchorte"]
    ergebnis: dict[str, list[int]] = {}
    for o in orte:
        s = plan.stufe(o)
        if o.ist_zusatz:
            schluessel = "Zusatz-Suchorte"
        elif o.einwohner is None:
            schluessel = "ohne Einwohnerzahl"
        else:
            schluessel = namen.get(id(s), "alle Orte")
        z = ergebnis.setdefault(schluessel, [0, 0, s.seiten or 0])
        z[0] += 1
        z[1] += len(plan.queries(o))
    return {k: tuple(ergebnis[k]) for k in sorted(ergebnis, key=lambda k: reihenfolge.index(k) if k in reihenfolge else 99)}
